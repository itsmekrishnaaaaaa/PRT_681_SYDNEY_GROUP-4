"""Assessment 3 deck.

Ink header, warm paper, copper rule. Facts match this codebase.
The ERD slide uses the final schema diagram, not the generated sketch.
"""

from __future__ import annotations

from pathlib import Path

from PIL import Image as PilImage
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parent
SHOT = ROOT / "shots"
OUTS = [
    ROOT / "OraEarn-Assessment3-Presentation.pptx",
    Path(r"C:\Users\budha\Downloads\OraEarn_Project_Presentation.pptx"),
    Path(r"C:\Users\budha\Downloads\S226 PRT681_Sydney_group4_Group_presentation.pptx"),
]

NAVY = RGBColor(0x14, 0x20, 0x2B)
NAVY_2 = RGBColor(0x1C, 0x2C, 0x3A)
INK = RGBColor(0x1C, 0x1A, 0x17)
MUTED = RGBColor(0x5E, 0x58, 0x50)
COPPER = RGBColor(0xB5, 0x68, 0x3D)
CREAM = RGBColor(0xF3, 0xE6, 0xD8)
PAPER = RGBColor(0xF4, 0xF0, 0xE6)
WHITE = RGBColor(0xFF, 0xFC, 0xF7)
LINE = RGBColor(0xD9, 0xD1, 0xC3)
ERD_SRC = Path(
    r"C:\Users\budha\.cursor\projects\c-projects-oraearn-pro\assets"
    r"\c__Users_budha_AppData_Roaming_Cursor_User_workspaceStorage_09925ccbc39933629b0361a31c7fee49_images_FinalERD.drawio__1_-4b6cf43b-0421-43be-977d-2e2fad6f90f0.png"
)

W = Inches(13.333)
H = Inches(7.5)


def _set_run(run, size, bold, color, font="Calibri"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font


def text(slide, l, t, w, h, body, size=18, bold=False, color=INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, font="Calibri"):
    box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = Emu(0)
    tf.margin_right = Emu(0)
    tf.margin_top = Emu(0)
    tf.margin_bottom = Emu(0)
    try:
        tf._txBody.bodyPr.set("anchor", {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}[anchor])
    except Exception:
        pass
    lines = body.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_before = Pt(0)
        p.space_after = Pt(0)
        run = p.add_run()
        run.text = line
        _set_run(run, size, bold, color, font)
    return box


def rect(slide, l, t, w, h, fill, line=None):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(1)
    sh.shadow.inherit = False
    return sh


def notes(slide, body):
    slide.notes_slide.notes_text_frame.text = body.strip()


def place_picture(slide, filename, l, t, w, h):
    path = SHOT / filename
    with PilImage.open(path) as im:
        iw, ih = im.size
    aspect = iw / ih
    if aspect > w / h:
        dw, dh = w, w / aspect
    else:
        dh, dw = h, h * aspect
    slide.shapes.add_picture(str(path), Inches(l + (w - dw) / 2), Inches(t + (h - dh) / 2), Inches(dw), Inches(dh))


def new(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
    bg.fill.solid()
    bg.fill.fore_color.rgb = PAPER
    bg.line.fill.background()
    sp_tree = slide.shapes._spTree
    sp = bg._element
    sp_tree.remove(sp)
    sp_tree.insert(2, sp)
    return slide


def header(slide, title, page):
    rect(slide, 0, 0, 13.333, 1.05, NAVY)
    rect(slide, 0, 1.05, 13.333, 0.045, COPPER)
    text(slide, 0.48, 0.14, 6, 0.22, "ORA EARN", size=11, bold=True, color=COPPER)
    text(slide, 0.48, 0.36, 10.8, 0.58, title, size=30, bold=True, color=WHITE, font="Georgia")
    text(slide, 11.15, 0.32, 1.7, 0.48, f"{int(page):02d}", size=22, bold=True, color=COPPER, align=PP_ALIGN.RIGHT, font="Georgia")
    text(slide, 0.48, 7.12, 8, 0.24, "PRT681    ·    Group 4", size=12, color=MUTED)
    text(slide, 9.3, 7.12, 3.55, 0.24, "Assessment 3", size=12, color=MUTED, align=PP_ALIGN.RIGHT)


def card(slide, l, t, w, h):
    rect(slide, l, t, w, h, WHITE, LINE)
    rect(slide, l, t, 0.07, h, COPPER)


def slide_title(prs):
    s = new(prs)
    rect(s, 0, 0, 13.333, 7.5, NAVY)
    rect(s, 0, 0, 13.333, 0.08, COPPER)
    rect(s, 0, 7.42, 13.333, 0.08, COPPER)
    text(s, 0.7, 0.45, 7.2, 0.28, "ASSESSMENT 3    ·    PRT681    ·    GROUP 4", size=13, bold=True, color=COPPER)
    text(s, 0.7, 1.05, 7.4, 1.15, "OraEarn", size=60, bold=True, color=WHITE, font="Georgia")
    text(s, 0.72, 2.35, 6.6, 0.9, "Earn smart. Stay compliant.\nSave with purpose.", size=22, color=CREAM, font="Georgia")
    rect(s, 0.72, 3.55, 1.4, 0.045, COPPER)
    text(
        s,
        0.7,
        6.35,
        6.8,
        0.7,
        "Our tutors and lecturers, for guidance and feedback,\nand every teammate whose work built OraEarn.",
        size=13,
        color=CREAM,
    )
    rect(s, 8.15, 0.7, 4.55, 5.85, NAVY_2)
    people = [
        ("Krishna Sapkota", "S395795", "Database administrator, tester"),
        ("Dharma Bahadur Budhathoki", "S395172", "Developer, business analyst"),
        ("Jenish Gurung", "S395321", "Business analyst, report writer"),
        ("Suraj Khadka", "S394949", "Data analyst, database administrator"),
    ]
    y = 0.95
    for name, sid, role in people:
        text(s, 8.4, y, 4.05, 0.55, name, size=15, bold=True, color=WHITE)
        text(s, 8.4, y + 0.48, 4.05, 0.24, sid, size=12, color=COPPER)
        text(s, 8.4, y + 0.72, 4.05, 0.4, role, size=13, color=CREAM)
        y += 1.35
    text(s, 11.5, 6.95, 1.3, 0.28, "01", size=14, bold=True, color=COPPER, align=PP_ALIGN.RIGHT, font="Georgia")
    notes(
        s,
        """
About 30 seconds. Read the line under the name, then the four people. Do not read the acknowledgements unless you are thanked.

Krishna: database and tests. Dharma: the build and the business view. Jenish: the rules and the written report. Suraj: the numbers and the tables.

Then: where the product came from.
""",
    )


def slide_origin(prs):
    s = new(prs)
    header(s, "Where we came from", 2)
    cols = [
        ("The challenge", "A student visa allows 48 hours in 14 days.\nShifts, expenses, and savings were tracked in separate places."),
        ("Our constraints", "First time on this stack.\nFive weekly sprints.\nFour people. We did not add a fifth late."),
        ("Our solution", "OraEarn. One account for jobs, shifts, the visa check, expenses, savings, and Aura."),
    ]
    x = 0.55
    for title, body in cols:
        card(s, x, 1.45, 3.9, 4.9)
        text(s, x + 0.28, 1.75, 3.35, 0.7, title, size=22, bold=True)
        text(s, x + 0.28, 2.7, 3.35, 3.1, body, size=18)
        x += 4.1
    notes(
        s,
        """
About 45 seconds. One sentence per column.

Challenge: the 48-hour fortnight, and money split across apps and spreadsheets.
Constraints: new technologies, five sprints from week 5 to week 9, four people. Brooks: adding a person late costs teaching time.
Solution: one login. The visa rule can see the shifts, because the shifts are in the same account.
""",
    )


def slide_product(prs):
    s = new(prs)
    header(s, "Product overview", 3)
    text(s, 0.55, 1.2, 12, 0.4, "Why OraEarn. One login covers the work, the limit, and the money.", size=18, color=MUTED)
    items = [
        ("Live shifts", "Clock in and out.\nPay comes from the job rate.\nThe calendar updates live."),
        ("Visa guard", "Rolling 14 days.\n48 hours for a student.\nOver the limit, the shift is not saved."),
        ("Money tools", "Expenses by category.\nSavings goals from leftover pay."),
        ("Aura and exports", "A small advisor beside the product.\nExcel and PDF from the same figures."),
    ]
    x = 0.55
    for title, body in items:
        card(s, x, 1.9, 2.95, 4.4)
        text(s, x + 0.22, 2.2, 2.5, 0.9, title, size=20, bold=True)
        text(s, x + 0.22, 3.3, 2.5, 2.5, body, size=16)
        x += 3.15
    notes(
        s,
        """
About 50 seconds. Point across the four cards.

Live shifts: SignalR hub /hubs/shifts. The hourly rate is on the job, not copied onto the shift.
Visa guard: VisaComplianceService. Student work rights only. Window is today minus 13 days through today. Limit is 48. The check runs before the insert.
Money: leftover pay is earnings minus expenses, never below zero.
Aura: Python on port 8000. If it is down, the rest of the app still runs. Exports are Excel and PDF of the same dashboard numbers.
""",
    )


def logo_tile(slide, l, t, w, h, filename, title, caption):
    card(slide, l, t, w, h)
    place_picture(slide, f"logos/{filename}", l + 0.2, t + 0.22, w - 0.4, 1.15)
    text(slide, l + 0.08, t + 1.45, w - 0.16, 0.32, title, size=14, bold=True, align=PP_ALIGN.CENTER)
    text(slide, l + 0.08, t + 1.78, w - 0.16, 0.45, caption, size=12, color=MUTED, align=PP_ALIGN.CENTER)


def slide_tech(prs):
    s = new(prs)
    header(s, "Technologies and tools", 4)
    logos = [
        ("react.png", "React 19", "The view. Port 5173."),
        ("typescript.png", "TypeScript", "The page code."),
        ("vite.png", "Vite", "The dev server."),
        ("dotnet.png", ".NET 10", "The API. Port 5203."),
        ("postgres.png", "PostgreSQL", "The database."),
        ("python.png", "Python", "Aura. Port 8000."),
    ]
    x = 0.4
    for filename, title, caption in logos:
        logo_tile(s, x, 1.22, 2.05, 2.4, filename, title, caption)
        x += 2.14
    text(
        s,
        0.55,
        3.75,
        12.2,
        0.35,
        "The page calls the API. The API calls PostgreSQL and, when needed, Aura.",
        size=15,
        color=MUTED,
    )
    card(s, 0.55, 4.25, 6.05, 2.5)
    text(s, 0.78, 4.45, 5.6, 0.35, "JWT", size=18, bold=True)
    text(
        s,
        0.78,
        4.9,
        5.6,
        1.6,
        "JwtBearer checks issuer, audience, lifetime, and the signing key.\n"
        "A good sign-in returns an access token and a refresh token.\n"
        "No token on a protected route returns 401.",
        size=14,
    )
    card(s, 6.8, 4.25, 5.95, 2.5)
    text(s, 7.05, 4.45, 5.5, 0.35, "Password is hashed", size=18, bold=True)
    text(
        s,
        7.05,
        4.9,
        5.5,
        1.6,
        "Identity writes PasswordHash. Sign-in compares the hash.\n"
        "At least 8 characters, with upper, lower, and a digit.\n"
        "Five failures lock the account for 15 minutes.",
        size=14,
    )
    notes(
        s,
        """
About 50 seconds. Walk the figure left to right, then the two cards.

React on 5173 is the view. It sends HTTP and the JWT to ASP.NET Core on 5203. Controllers call services. Services use EF Core. EF Core talks to PostgreSQL. The browser never opens the database.

SignalR hub is /hubs/shifts and uses the same token. Aura is Python on port 8000. The API calls it. The product still runs if Aura is off.

JWT: AddJwtBearer. Checks issuer, audience, lifetime, and the signing key. Access token plus refresh token. Logout revokes the refresh token.

Password: UserManager.CreateAsync stores PasswordHash. CheckPasswordSignInAsync compares the hash. Length 8, digit, lower, upper. Lockout is 5 failures, then 15 minutes.
""",
    )


def slide_architecture(prs):
    s = new(prs)
    header(s, "Architecture", 5)
    place_picture(s, "diagram-architecture.png", 0.35, 1.12, 12.6, 5.75)
    notes(
        s,
        """
About 60 seconds. Three layers, left to right, then the return line, then Aura.

Client: the React pages on port 5173. There is no Razor page. The browser does not open PostgreSQL.
Presentation layer: AuthController, JobsController, ShiftsController, ExpensesController, SavingsGoalsController, AnalyticsController, NotificationsController, AdminController. A protected action has Authorize. The JWT is checked before the method runs.
Business logic layer: the controller calls a service. ShiftService, JobService, AuthService, ExpenseService, VisaComplianceService. The 48-hour check is in the service, before any insert.
Data access layer: ApplicationDbContext and EF Core. PostgreSQL is behind that. There is no repository class.

The copper line is the JSON response, back up through each layer to the same page.

Aura hangs off the business side. PythonAdvisorClient calls port 8000. The page does not call Python. Aura does not open the database.

SignalR is on the same API. The hub is authorized with the same JWT.
""",
    )


def figure_slide(prs, title, page, filename, spoken, caption=None):
    s = new(prs)
    header(s, title, page)
    top = 1.28
    if caption:
        text(s, 0.5, 1.22, 12.3, 0.32, caption, size=15, color=MUTED)
        top = 1.58
    place_picture(s, filename, 0.35, top, 12.6, 6.95 - top)
    notes(s, spoken)
    return s


def slide_data(prs):
    s = new(prs)
    header(s, "Database design and integrity", 6)
    text(s, 0.55, 1.18, 12.2, 0.35, "Third normal form. Each fact is stored once.", size=16, color=MUTED)

    def box(l, t, w, label):
        rect(s, l, t, w, 0.7, WHITE, INK)
        text(s, l, t + 0.16, w, 0.4, label, size=16, bold=True, align=PP_ALIGN.CENTER)

    box(0.55, 1.75, 2.3, "Users")
    text(s, 2.9, 1.9, 0.7, 0.4, "1 → *", size=14, color=MUTED, align=PP_ALIGN.CENTER)
    box(3.6, 1.75, 2.5, "Jobs")
    text(s, 6.15, 1.9, 0.7, 0.4, "1 → *", size=14, color=MUTED, align=PP_ALIGN.CENTER)
    box(6.9, 1.75, 2.4, "Shifts")
    text(s, 9.5, 1.85, 3.3, 0.55, "Rate stays on the job.", size=14, color=MUTED)

    box(0.55, 2.85, 2.3, "Users")
    text(s, 2.85, 3.0, 0.75, 0.4, "1 → *", size=14, color=MUTED, align=PP_ALIGN.CENTER)
    box(3.6, 2.85, 2.7, "Categories")
    text(s, 6.35, 3.0, 0.7, 0.4, "1 → *", size=14, color=MUTED, align=PP_ALIGN.CENTER)
    box(7.05, 2.85, 2.3, "Expenses")
    text(s, 9.4, 3.0, 0.85, 0.4, "* ← 1", size=14, color=MUTED, align=PP_ALIGN.CENTER)
    box(10.25, 2.85, 2.5, "Pay methods")

    box(0.55, 3.95, 1.9, "Users")
    text(s, 2.45, 4.1, 0.7, 0.4, "1 → *", size=14, color=MUTED, align=PP_ALIGN.CENTER)
    box(3.15, 3.95, 2.55, "Savings goals")
    box(6.15, 3.95, 1.9, "Users")
    text(s, 8.05, 4.1, 0.9, 0.4, "1 → 0..1", size=14, color=MUTED, align=PP_ALIGN.CENTER)
    box(8.95, 3.95, 2.5, "Live shift")

    card(s, 0.55, 5.05, 12.2, 1.7)
    text(
        s,
        0.8,
        5.28,
        11.7,
        1.3,
        "An expense stores a category and a payment method. A shift stores the job. The hourly rate stays on the job.\n"
        "Roles are their own table, joined through User_Roles. A notification read is one row per user.\n"
        "48 hours is a sum of shift hours over the last 14 days. It is not a column.",
        size=16,
    )
    notes(
        s,
        """
About 60 seconds. Walk the three lines, then the sentence at the bottom.

Line 1: a user has jobs. A job has shifts. The shift stores JobId, not a copied hourly rate, and not a UserId. Pay is calculated from the job's rate at save time. Editing the job rate later does not rewrite old shifts.

Line 2: this is the correction if someone remembers an older diagram. An expense does not link to a job. It links to ExpenseCategory and PaymentMethod. The category belongs to the user, so the expense is still that user's money. Payment methods are Card, Cash, Bank Transfer, and Other. Delete is Restrict while expenses exist.

Line 3: savings goals belong to the user. Allocation is a percent of leftover pay. The live shift is ActiveShiftSession. Its UserId is unique, so only one clock-in is open.

The migration is named NormalizeTo3NF. The 48-hour cap is VisaComplianceService, a query, not a column.

Roles are not a column on the user. User_Roles joins a user to a role. User_Notification_Reads stores which notification that user has opened. Both are on the final ERD.
""",
    )


def slide_security(prs):
    s = new(prs)
    header(s, "Security and quality assurance", 11)
    steps = [
        ("1  Sign in", "Identity stores PasswordHash.\nThe typed password is not saved.\nA good sign-in returns a JWT."),
        ("2  Approval", "New accounts are Pending.\nPending and Rejected get no token.\nAn admin approves first."),
        ("3  Role", "Most routes need a token.\nAdmin routes also need\nthe Admin role."),
        ("4  Own data", "A query is limited to that user.\nOne student cannot read\nanother student's pay."),
    ]
    x = 0.55
    for title, body in steps:
        card(s, x, 1.25, 2.95, 2.55)
        text(s, x + 0.16, 1.42, 2.65, 0.4, title, size=16, bold=True)
        text(s, x + 0.16, 1.9, 2.65, 1.7, body, size=13)
        x += 3.15
    card(s, 0.55, 4.05, 6.0, 2.7)
    text(s, 0.8, 4.25, 5.5, 0.4, "Testing", size=18, bold=True)
    text(
        s,
        0.8,
        4.75,
        5.5,
        1.8,
        "29 tests passed. None failed. None skipped.\n"
        "Bad login returns 401.\n"
        "No token on /me, or on jobs, returns 401.\n"
        "A valid token returns the profile.",
        size=15,
    )
    card(s, 6.75, 4.05, 6.0, 2.7)
    text(s, 7.0, 4.25, 5.5, 0.4, "What we still owe", size=18, bold=True)
    text(
        s,
        7.0,
        4.75,
        5.5,
        1.8,
        "Repeated bad passwords lock the account.\n"
        "Logout revokes the refresh token.\n"
        "Gap: no test yet posts a shift that\nwould cross 48 hours.",
        size=15,
    )
    notes(
        s,
        """
About 70 seconds. This is the security story. Do not claim a hash algorithm by name. Say ASP.NET Identity writes PasswordHash when the user is created. CheckPasswordSignInAsync compares that hash. Lockout is on.

Approval: Register sets ApprovalStatus to Pending. A non-admin who is Pending or Rejected never reaches a token. Admin can sign in to do the approval.

Unauthorised tests, black box: Login_WithInvalidCredentials returns 401. Me without a token returns 401. GetJobs without a token returns 401. A good login returns an access token. Me with that token returns the profile.

Be straight about the gap. The 48-hour refusal is in the service. The test project does not yet post a shift that crosses 48 and expect a rejection. That is the next test. Alice on the live screen is at 7.5 of 48, so do not say she is already over.
""",
    )


def slide_analytics(prs):
    s = new(prs)
    header(s, "Data, analytics, and Aura", 12)
    cols = [
        ("Shifts and expenses", "Hours, income, and spending.\nPay is hours times the job rate, after the break."),
        ("Analytics", "Same formula.\nWindows of 7 days, 30 days,\n1 month, and 12 months."),
        ("Goal tracking", "Leftover pay is earnings minus expenses, never below zero.\nProgress caps at 100 percent."),
        ("Aura", "Answers a savings question from that user's figures.\nIf Aura is off, the app still runs."),
    ]
    x = 0.55
    for title, body in cols:
        card(s, x, 1.4, 2.95, 5.15)
        text(s, x + 0.2, 1.7, 2.55, 1.1, title, size=20, bold=True)
        text(s, x + 0.2, 3.0, 2.55, 3.1, body, size=16)
        x += 3.15
    notes(
        s,
        """
About 50 seconds. Suraj owns these numbers if he is speaking. Use only the live screen.

Alice, student, last 7 days, 29 Sep 2026 to 5 Oct 2026: one shift, 7.5 hours, earnings $210, expenses $365, net −$155. 7.5 times $28 is $210, so that shift was priced from Cafe Barista. Visa bar: 7.5 of 48, 40.5 left, window 22 Sep 2026 to 5 Oct 2026. Savings on screen: targets about $9,722, 65 percent allocated, about $1,144 saved. Point at Holiday trip and Emergency fund. Ignore the short test names.

Two sums: the 7-day card is not the savings total. Goals use all earnings minus all expenses.

Hours still needed is remaining dollars divided by the hourly rate. If allocation percents pass 100, they are scaled back to 100.
""",
    )


def slide_process(prs):
    s = new(prs)
    header(s, "Development process", 13)
    weeks = [
        ("Week 5", "Foundation", "Account and database."),
        ("Week 6", "Shifts", "Jobs, clock-in, SignalR."),
        ("Week 7", "Money", "Expenses and goals."),
        ("Week 8", "Compliance", "48-hour rule, analytics, exports."),
        ("Week 9", "Aura", "Notifications and Aura."),
    ]
    x = 0.55
    for week, title, body in weeks:
        card(s, x, 1.25, 2.35, 2.7)
        text(s, x + 0.14, 1.42, 2.05, 0.3, week, size=13, color=MUTED)
        text(s, x + 0.14, 1.75, 2.05, 0.55, title, size=18, bold=True)
        text(s, x + 0.14, 2.4, 2.05, 1.3, body, size=14)
        x += 2.52
    bottom = [
        ("Agile / Scrum", "Five weekly sprints, weeks 5 to 9.\nA sprint goal, then the stories.\nDone when someone can use it."),
        ("Grooming", "Login, then shifts, then the 48-hour rule, then Aura.\nIf the week slipped, the next goal changed."),
        ("Mythical man-month", "Four people, week 5 through week 9.\nWe did not add a fifth.\nA new person costs teaching time."),
    ]
    x = 0.55
    for title, body in bottom:
        card(s, x, 4.15, 3.95, 2.55)
        text(s, x + 0.22, 4.35, 3.5, 0.4, title, size=18, bold=True)
        text(s, x + 0.22, 4.9, 3.5, 1.5, body, size=15)
        x += 4.15
    notes(
        s,
        """
About 40 seconds. This slide is the process. The next slide is the practices.

Point across the weeks. The order is the grooming. The visa check is a sum of shifts, so shifts existed before the rule. Aura has nothing to say until goals and hours exist. That is the iteration: each week is a sprint with a goal, and a story is done when someone can use it.

Mythical man-month: we stayed four people. Adding someone in week 8 or 9 would have been spent on teaching. A week moved a story. It did not grow the team.
""",
    )


def slide_practices(prs):
    s = new(prs)
    header(s, "Practices and the team", 14)
    text(
        s,
        0.55,
        1.2,
        12.2,
        0.32,
        "Git, task names, an equal share, and why a customer would keep this team.",
        size=15,
        color=MUTED,
    )
    cols = [
        (
            "Git history",
            "Form we keep\n"
            "feat: what changed\n"
            "fix: what broke\n\n"
            "The log has both.\n"
            "Early rows also say\n"
            "Add files via upload.\n"
            "That form stops.",
        ),
        (
            "Task names",
            "OE-101 Register account\n"
            "As a / I want / So that\n"
            "Acceptance criteria\n"
            "on the story.\n\n"
            "Board: Backlog, To Do,\n"
            "In Progress, In Review,\n"
            "UAT, Done.",
        ),
        (
            "Equal share",
            "Krishna: database, tests\n"
            "Dharma: the build\n"
            "Jenish: rules, report\n"
            "Suraj: the numbers\n\n"
            "Each person led stories.\n"
            "The others took sub-tasks.\n"
            "The plan gave the same hours.",
        ),
        (
            "Keep this team",
            "Yes.\n\n"
            "The sprint changed.\n"
            "The headcount did not.\n"
            "Next work is the\n"
            "48-hour test.\n"
            "This team already\n"
            "holds the code.",
        ),
    ]
    x = 0.4
    for title, body in cols:
        card(s, x, 1.62, 3.15, 5.15)
        text(s, x + 0.2, 1.8, 2.8, 0.7, title, size=18, bold=True)
        text(s, x + 0.2, 2.55, 2.8, 3.9, body, size=14)
        x += 3.25
    notes(
        s,
        """
About 70 seconds. Four cards. This is the development-process answer.

Git: point at the form, then be honest. Working messages are feat: and fix:, for example feat: added the password toggle, and fix: for the expense screens. The history also contains "Add files via upload" from the GitHub website. Say that those are in the log, and the form we keep is type, colon, what changed. Do not invent a clean history.

Tasks: the name is the key plus the title, OE-101 Register account. The body is As a, I want, So that, then acceptance criteria. Grooming ordered them by dependency, which is the week row on the previous slide. Prioritisation was that order, not whoever spoke last. The board is Backlog, To Do, In Progress, In Review, UAT, Done.

Equal share: Krishna held the database and the tests. Dharma held the build. Jenish held the rules and the written report. Suraj held the numbers. Each person led stories. The other three took sub-tasks on those stories, so one person did not hold the week. The sprint plan gave every person the same hour share. Do not quote a timesheet total.

Adaptive: when a dependency was missing, the story moved to the next sprint. The team size stayed four. That is the evolutionary practice. We changed the plan. We did not change the headcount. That is also the mythical man-month: a fifth person in the last weeks would have been teaching time.

Customer: the answer is yes. The product is already in slices someone can use. The next job is known, the test that expects the 48-hour refusal. This team already holds the code, so continued development does not start with onboarding.
""",
    )


def slide_future(prs):
    s = new(prs)
    header(s, "Strengths, improvements, and future", 15)
    cols = [
        ("Key strengths", "The visa rule blocks the save.\nThe database is in third normal form.\nA clock-out updates other open tabs.\nThe demo data can be seeded again.\n29 tests are green."),
        ("Improve next", "A test that posts a shift past 48 hours.\nMore tests around approval and lockout.\nA pass over slow dashboard queries."),
        ("Future features", "A phone app on the same API.\nOther visa types, not only the student 48.\nA bank feed, so expenses are not typed.\nA deeper Aura, same panel."),
    ]
    x = 0.55
    for title, body in cols:
        card(s, x, 1.4, 3.95, 5.2)
        text(s, x + 0.25, 1.7, 3.45, 0.7, title, size=22, bold=True)
        text(s, x + 0.25, 2.6, 3.45, 3.6, body, size=16)
        x += 4.15
    notes(
        s,
        """
About 40 seconds. Left is what works today. Middle is the honest gap. Right is later, not this assessment.

The missing test matters because the rule is real in VisaComplianceService and the screen shows Alice at 7.5 of 48. The suite proves hour maths, pay, savings, and 401s. It does not yet prove the refusal.

Future: the API is already there for a phone client. Aura stays a separate Python service.
""",
    )


def slide_demo(prs):
    s = new(prs)
    header(s, "Live demo", 16)
    text(s, 0.55, 1.18, 12, 0.35, "OraEarn in action. Alice is the student. Bob is not on the 48-hour rule.", size=16, color=MUTED)
    cols = [
        ("Shifts", "Dharma", "Sign in as Alice.\nAdd a shift.\nThe calendar updates on the open tab."),
        ("Money", "Dharma", "Add an expense.\nAdd a savings goal.\nLeftover pay moves the goal."),
        ("Visa guard", "Krishna", "A shift that would pass 48 hours is refused.\nBob, a citizen, is not on that rule."),
        ("Aura and export", "Krishna", "Open analytics.\nExport.\nAsk Aura.\nThen questions."),
    ]
    x = 0.55
    for title, who, body in cols:
        card(s, x, 1.75, 2.95, 4.85)
        text(s, x + 0.2, 2.0, 2.55, 0.7, title, size=20, bold=True)
        text(s, x + 0.2, 2.75, 2.55, 0.4, who, size=14, bold=True, color=MUTED)
        text(s, x + 0.2, 3.3, 2.55, 2.9, body, size=15)
        x += 3.15
    notes(
        s,
        """
Demo accounts. Do not read the passwords to the room unless the marker asks you to sign in.

Alice Nguyen, student, already approved: alice@oraearn.local / User123!
Bob, citizen, full work rights: bob@oraearn.local / User123!
Admin: admin@oraearn.local / Admin123!

Alice's visa bar is 7.5 of 48, window 22 Sep 2026 to 5 Oct 2026. She is not over the limit. To show the guard, enter a shift that would take the fortnight over 48. The API refuses and does not insert. Then sign in as Bob and show that the same rule does not apply.

Ports: frontend 5173, API 5203, Aura 8000. Swagger is /swagger if they ask for the lock on each route.

Jobs to point at: Cafe Barista $28, Retail Assistant $26. Ignore a test job named dfdf. Savings: Holiday trip and Emergency fund. Ignore a test goal named asd.

Close on questions. The next test we would write is the one that expects the 48-hour refusal.
""",
    )


def install_final_erd():
    local = SHOT / "erd-final-source.png"
    src = ERD_SRC if ERD_SRC.exists() else local
    if not src.exists():
        raise FileNotFoundError(f"Final ERD image not found: {src}")
    with PilImage.open(src) as im:
        rgb = im.convert("RGB")
        rgb.save(local, "PNG")
        rgb.save(SHOT / "diagram-erd.png", "PNG")


def build():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    prs.core_properties.title = "OraEarn — Assessment 3"
    prs.core_properties.subject = "PRT681 Group 4"
    prs.core_properties.category = "Assessment 3"
    import draw_figures

    draw_figures.main()
    install_final_erd()
    slide_title(prs)
    slide_origin(prs)
    slide_product(prs)
    slide_tech(prs)
    slide_architecture(prs)
    slide_data(prs)
    figure_slide(
        prs,
        "Entity relationship diagram",
        7,
        "diagram-erd.png",
        """
About 50 seconds. This is the final ERD. Point around the page, do not read every column.

Roles connect to users through User_Roles. A role has a name and a normalized name. The join table holds UserId and RoleId.

Users hold the email, username, visa fields, the identity document, approval, and the refresh token. A user has many jobs. A job holds the hourly rate and an optional end date. A shift stores JobId, the times, the break, total hours, and the earned snapshot. The rate is not copied onto the shift as the job's rate column.

Active shift sessions are the open clock-in. UserId there is unique, so only one is open.

Expense categories belong to the user. An expense stores CategoryId and PaymentMethodId, the amount, the date, and fixed or variable. Payment method names are unique.

Savings goals belong to the user: target, current amount, and the allocation percent. User notification reads are one row for that user and that notification.

48 hours is a sum of shift hours over 14 days. It is not a table on this diagram.
""",
        caption="Final schema. Eleven tables. The hourly rate stays on the job.",
    )
    figure_slide(
        prs,
        "Normalization",
        8,
        "diagram-normalization.png",
        """
About 40 seconds. Walk left to right.

UNF was one spreadsheet row: visa, jobs, shifts, and expenses packed together.
1NF made each fact a row, but the shift still copied the user, the job title, and the rate.
3NF is what is running. The rate lives on the job. The shift stores JobId. The expense stores two foreign keys.

The migration is named NormalizeTo3NF. A shift already saved keeps the pay it was given when Cafe's rate changes.
""",
    )
    figure_slide(
        prs,
        "Use case",
        9,
        "diagram-usecase.png",
        """
About 35 seconds. Two actors.

The student registers, signs in, works with jobs and shifts, watches the 48-hour bar, logs money, exports, and can ask Aura.
The admin approves or rejects, opens the ID file, sets the role, and can see platform analytics.

A student action and an admin action are separate. Admin routes need the Admin role.
""",
    )
    figure_slide(
        prs,
        "Sequence: saving a shift",
        10,
        "diagram-sequence.png",
        """
About 50 seconds. Follow the numbers.

1. React posts /api/Shifts with the JWT.
2. Shift.Calculate uses the job's HourlyRate. Hours are end minus start minus break.
3. VisaCompliance runs only for a student. It sums hours over the last 14 days.
4 and 5. That sum is a query on this user's shifts.
6. Over 48, the save stops. Nothing is inserted.
7. Under 48, the shift row is inserted.
8. The API returns 201 and the savings goals are recalculated.

The check is in ShiftService.CreateAsync, before SaveChanges.
""",
    )
    slide_security(prs)
    slide_analytics(prs)
    slide_process(prs)
    slide_practices(prs)
    slide_future(prs)
    slide_demo(prs)
    for path in OUTS:
        path.parent.mkdir(parents=True, exist_ok=True)
        try:
            prs.save(path)
            print(f"Saved {path}")
        except PermissionError:
            alt = path.with_name(path.stem + "-ink.pptx")
            prs.save(alt)
            print(f"Locked, saved {alt}")
    print(f"Slides: {len(prs.slides)}")


if __name__ == "__main__":
    build()
