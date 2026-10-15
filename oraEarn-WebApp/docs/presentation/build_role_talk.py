"""Separate talk for Dharma: developer and Scrum Master.

He speaks first about what he did, then asks the team.
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parent
OUTS = [
    ROOT / "OraEarn-Dharma-Role-Questions.pptx",
    Path(r"C:\Users\budha\Downloads\OraEarn-Dharma-Role-Questions.pptx"),
]

INK = RGBColor(0x11, 0x11, 0x11)
MUTED = RGBColor(0x44, 0x44, 0x44)
HAIR = RGBColor(0xC5, 0xD5, 0xD4)
SOFT = RGBColor(0xE7, 0xF4, 0xF3)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
TEAL = RGBColor(0x0E, 0x7C, 0x7B)

W = Inches(13.333)
H = Inches(7.5)


def _set_run(run, size, bold, color, font="Calibri"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font


def text(slide, l, t, w, h, body, size=18, bold=False, color=INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = Emu(0)
    tf.margin_right = Emu(0)
    tf.margin_top = Emu(0)
    tf.margin_bottom = Emu(0)
    try:
        tf._txBody.bodyPr.set("anchor", {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}[anchor])
    except Exception:
        pass
    lines = body.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_before = Pt(0)
        p.space_after = Pt(0)
        run = p.add_run()
        run.text = line
        _set_run(run, size, bold, color)
    return box


def rect(slide, l, t, w, h, fill, line=None):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(1)
    sh.shadow.inherit = False
    return sh


def notes(slide, body):
    slide.notes_slide.notes_text_frame.text = body.strip()


def new(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
    bg.fill.solid()
    bg.fill.fore_color.rgb = WHITE
    bg.line.fill.background()
    sp_tree = slide.shapes._spTree
    sp = bg._element
    sp_tree.remove(sp)
    sp_tree.insert(2, sp)
    return slide


def card(slide, l, t, w, h):
    rect(slide, l, t, w, h, SOFT, HAIR)
    rect(slide, l, t, w, 0.07, TEAL)


def footer(slide, page):
    text(slide, 11.4, 7.05, 1.4, 0.28, str(page), size=14, bold=True, color=TEAL, align=PP_ALIGN.RIGHT)
    text(slide, 0.55, 7.05, 9, 0.28, "OraEarn   ·   Dharma   ·   Developer and Scrum Master", size=12, color=MUTED)


def slide_title(prs):
    s = new(prs)
    rect(s, 0, 0, 13.333, 0.1, TEAL)
    rect(s, 0, 0, 0.12, 7.5, TEAL)
    text(s, 0.6, 0.4, 12, 0.35, "PRT681    ·    GROUP 4    ·    I SPEAK FIRST", size=14, bold=True, color=TEAL)
    text(s, 0.6, 0.9, 12, 0.8, "Dharma Bahadur Budhathoki", size=36, bold=True)
    text(s, 0.6, 1.75, 12, 0.4, "S395172", size=18, color=MUTED)
    text(s, 0.6, 2.3, 12, 0.7, "Developer and Scrum Master", size=28, bold=True)
    card(s, 0.6, 3.3, 5.8, 2.3)
    text(s, 0.85, 3.55, 5.3, 0.4, "First I tell you what I did", size=18, bold=True)
    text(s, 0.85, 4.1, 5.3, 1.2, "Scrum Master: the order of the work,\nand why we stayed four.\nDeveloper: the API, the hash, the visa check.", size=16)
    card(s, 6.7, 3.3, 5.9, 2.3)
    text(s, 6.95, 3.55, 5.4, 0.4, "Then I ask everyone", size=18, bold=True)
    text(s, 6.95, 4.1, 5.4, 1.2, "Five technical questions.\nFive ethical questions.\nI answer first. Then Jenish, Krishna, and Suraj.", size=16)
    footer(s, 1)
    notes(
        s,
        """
About 20 seconds.

I am Dharma. I speak first because I am the Scrum Master and the developer. I will tell you what I did in those two roles. After that I will ask five technical questions and five ethical questions. I answer each one first, then I ask Jenish, Krishna, and Suraj to add their part.
""",
    )


def slide_scrum(prs):
    s = new(prs)
    rect(s, 0, 0, 0.12, 7.5, TEAL)
    text(s, 0.55, 0.28, 12, 0.3, "SCRUM MASTER", size=14, bold=True, color=TEAL)
    text(s, 0.55, 0.62, 12, 0.5, "What I did first", size=32, bold=True)
    rect(s, 0.55, 1.2, 12.23, 0.035, TEAL)
    points = [
        ("The idea and the order", "I owned the product idea, and I set the order the team built it in."),
        ("Five phases", "Account, then jobs and shifts, then money, then the 48-hour rule, then Aura."),
        ("Four people", "Jenish on the rules, Krishna on the tables, Suraj on the numbers. I did not add a fifth person."),
        ("Done meant usable", "A phase was finished when you could use it, not when a document said so."),
    ]
    y = 1.5
    for title, body in points:
        card(s, 0.55, y, 12.2, 1.2)
        text(s, 0.8, y + 0.18, 11.7, 0.35, title, size=18, bold=True)
        text(s, 0.8, y + 0.58, 11.7, 0.45, body, size=16)
        y += 1.32
    footer(s, 2)
    notes(
        s,
        """
About 50 seconds. Say this before any question.

I am the Scrum Master, so I speak before the others. The idea is mine, and so is the order. Account had to work first, because nothing else is safe until a person can sign in. Jobs and shifts next, because the pay rate lives on the job. Money after that. The 48-hour rule is a sum of shifts, so it waited until shifts existed. Aura was last, because it needs hours and goals before it has anything to say.

We stayed four. If we had slipped, I would have moved a story. I would not have opened a fifth seat in the last weeks. Jenish, Krishna, and Suraj each had a clear part. They will go deeper after I ask them.
""",
    )


def slide_developer(prs):
    s = new(prs)
    rect(s, 0, 0, 0.12, 7.5, TEAL)
    text(s, 0.55, 0.28, 12, 0.3, "DEVELOPER", size=14, bold=True, color=TEAL)
    text(s, 0.55, 0.62, 12, 0.5, "What I built", size=32, bold=True)
    rect(s, 0.55, 1.2, 12.23, 0.035, TEAL)
    points = [
        ("The page talks only to the API", "React on 5173. ASP.NET Core on 5203. The browser never opens PostgreSQL."),
        ("Sign-in", "Identity stores PasswordHash. A good sign-in returns a JWT. Pending accounts get no token."),
        ("Services hold the rules", "The controller calls a service. The 48-hour check runs before a shift is inserted."),
        ("Tests, and the one I still owe", "29 passed, none failed. There is still no test that posts a shift past 48 hours."),
    ]
    y = 1.5
    for title, body in points:
        card(s, 0.55, y, 12.2, 1.2)
        text(s, 0.8, y + 0.18, 11.7, 0.35, title, size=18, bold=True)
        text(s, 0.8, y + 0.58, 11.7, 0.45, body, size=16)
        y += 1.32
    footer(s, 3)
    notes(
        s,
        """
About 50 seconds. Still me, now as the developer.

The React page only calls the API. JWT is checked on the API. Register stores PasswordHash through ASP.NET Identity. A Pending or Rejected account does not receive a token.

Controllers stay thin. ShiftService loads the job rate, calculates the hours, and asks VisaComplianceService before it inserts. Over 48, nothing is stored.

I ran the tests for this talk. 29 passed. The missing one is a shift that would cross 48 hours. Alice on the live screen is at 7.5 of 48. I will not say she is already over.
""",
    )


def slide_ask(prs):
    s = new(prs)
    rect(s, 0, 0, 0.12, 7.5, TEAL)
    text(s, 0.55, 0.28, 12, 0.3, "THEN I ASK", size=14, bold=True, color=TEAL)
    text(s, 0.55, 0.62, 12, 0.6, "I answer first. Then I ask all three.", size=28, bold=True)
    rect(s, 0.55, 1.4, 12.23, 0.035, TEAL)
    people = [
        ("Jenish Gurung", "The rules", "Approval, who the 48 hours apply to, and what may be saved."),
        ("Krishna Sapkota", "The database and the tests", "Where each fact lives, and the test we still owe."),
        ("Suraj Khadka", "The numbers", "Alice’s 7.5 hours, the $210, and the two different sums."),
    ]
    x = 0.55
    for name, role, body in people:
        card(s, x, 1.75, 3.95, 3.4)
        text(s, x + 0.22, 2.0, 3.5, 0.8, name, size=20, bold=True)
        text(s, x + 0.22, 2.85, 3.5, 0.4, role, size=15, bold=True, color=TEAL)
        text(s, x + 0.22, 3.4, 3.5, 1.4, body, size=16)
        x += 4.15
    text(s, 0.55, 5.4, 12.2, 1.2, "Five technical questions, then five ethical ones.\nI answer first, then I ask Jenish, Krishna, and Suraj.", size=18)
    footer(s, 4)
    notes(s, "Do not linger. Start technical question 1. I answer, then I name who I am asking.")


def question_slide(prs, section, question, mine, ask, page):
    s = new(prs)
    rect(s, 0, 0, 0.12, 7.5, TEAL)
    text(s, 0.55, 0.22, 12, 0.28, section, size=13, bold=True, color=TEAL)
    text(s, 0.55, 0.52, 12.2, 1.05, question, size=22, bold=True)
    card(s, 0.55, 1.62, 12.2, 2.15)
    text(s, 0.8, 1.76, 11.7, 0.26, "I say this first", size=14, bold=True, color=TEAL)
    text(s, 0.8, 2.06, 11.7, 1.55, mine, size=15)
    card(s, 0.55, 3.92, 12.2, 2.85)
    text(s, 0.8, 4.06, 11.7, 0.26, "Then I ask the others", size=14, bold=True, color=TEAL)
    text(s, 0.8, 4.38, 11.7, 2.2, ask, size=14)
    footer(s, page)
    notes(s, mine + "\n\nThen I ask: " + ask)
    return s


def build():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    prs.core_properties.title = "OraEarn — Dharma, developer and Scrum Master"
    prs.core_properties.subject = "PRT681 Group 4"
    slide_title(prs)
    slide_scrum(prs)
    slide_developer(prs)
    slide_ask(prs)
    questions = [
        (
            "TECHNICAL  ·  1 OF 5",
            "GET shows the data in the URL, and it can post as well. PUT and PATCH are the same, right?",
            "No. A GET puts the query in the URL, so it lands in logs and history. A POST sends the body, and POST /api/Shifts inserts a row. We have no PATCH. PUT /api/Shifts/{id} replaces the date, the times, the break, and the job, then recalculates the pay. PATCH would change only the field you sent. A short PUT can wipe the fields you left out.",
            "Jenish: if Save is a GET, is the shift stored, and is the body visible in the address bar?\nKrishna: we send a full ShiftRequest on PUT. What happens to EndTime if the body omits it?\nSuraj: POST the same $210 shift twice. Does the week show $210 or $420?",
            5,
        ),
        (
            "TECHNICAL  ·  2 OF 5",
            "Why localhost, and is the refresh token the same as the access token?",
            "Localhost is this machine only: page 5173, API 5203, Aura 8000. A cookie rides with the browser. A session sits in server memory. We use two tokens. The access token is a signed JWT for 15 minutes. It goes on every call, and it is not stored. The refresh token is random, stored on the user for 7 days, and it only asks for a new access token. Logout clears it. An expired access token does not stay valid.",
            "Jenish: a Pending login is stopped before the hash is checked. Is that still 401?\nKrishna: the socket reads the access token from the query. If that 15-minute token is expired, does the 7-day refresh token keep it open?\nSuraj: can the dashboard dollars be read with the refresh token alone?",
            6,
        ),
        (
            "TECHNICAL  ·  3 OF 5",
            "What is an ORM, and what does it actually give us?",
            "Ours is Entity Framework Core. The classes become the tables. It writes the SQL, tracks what changed, and runs the migrations, including NormalizeTo3NF. It also gives us the foreign keys and the delete rules. It does not know the 48-hour rule. That stays in the visa service. A junior thinks the ORM is the business rule. It is only the mapping.",
            "Jenish: if EF can save a shift, why does ShiftService still call the visa check first?\nKrishna: which delete rule did EF map, Cascade or Restrict, when a job still has shifts?\nSuraj: does the ORM store the 48-hour cap, or is that number calculated when the card is drawn?",
            7,
        ),
        (
            "TECHNICAL  ·  4 OF 5",
            "A protected route in React is enough, right?",
            "No. ProtectedRoute only sends a signed-out person to /login, and AdminRoute guards the admin page. Hiding a page is not security. The API has [Authorize]. No token is 401. A student token on an admin action fails the role check. That is 403. A junior protects the route and leaves the API open. Then Swagger still saves the shift.",
            "Jenish: if I delete ProtectedRoute but leave [Authorize], can a stranger still POST a shift?\nKrishna: a student token calls approve. Is that 401 or 403?\nSuraj: if the page is hidden but the API is open, can someone still read Alice’s $210?",
            8,
        ),
        (
            "TECHNICAL  ·  5 OF 5",
            "What fills the server, and how do we keep it small?",
            "The shift and expense rows are small. The heavy files are the ID uploads in uploads/visas, and each one is capped at 5 MB, PDF or image. node_modules, bin, and obj are huge, and they are not what we deploy. We do not delete old shifts to save disk. Those rows are the pay record. We publish the built app, and we keep the visa files limited.",
            "Jenish: is it acceptable to delete last month’s shifts so the disk stays small?\nKrishna: the database stores the visa path, not the file bytes. Where do the bytes actually sit?\nSuraj: if two ID files are saved, does Alice’s $210 card get bigger, or only the disk?",
            9,
        ),
        (
            "ETHICAL  ·  1 OF 5",
            "Who can open the ID file?",
            "Only an admin. The file is there so they can approve the person. A new account stays Pending until they do.",
            "Jenish: can a normal user open it?\nKrishna: is the file stored in the password table?\nSuraj: would you put it in an export?",
            10,
        ),
        (
            "ETHICAL  ·  2 OF 5",
            "Is it all right to print a password while we debug?",
            "No. We store PasswordHash. The typed password stays out of the log, the table, and the API response.",
            "Jenish: would you print it to debug?\nKrishna: is the typed password in the table?\nSuraj: would you put it in a spreadsheet?",
            11,
        ),
        (
            "ETHICAL  ·  3 OF 5",
            "Why is Alice limited, and Bob is not?",
            "Alice is a student, so the limit is 48 hours in 14 days. Bob is a citizen, so that limit is not applied to him.",
            "Jenish: who is limited, Alice or Bob?\nKrishna: is 48 a column in the database?\nSuraj: what does Alice’s bar show today?",
            12,
        ),
        (
            "ETHICAL  ·  4 OF 5",
            "Can one student see another student’s pay?",
            "No. The query uses the signed-in user’s id. Pay, visa type, and savings stay private.",
            "Jenish: is one student allowed to open another’s pay?\nKrishna: is that pay private?\nSuraj: would an export include someone else’s rows?",
            13,
        ),
        (
            "ETHICAL  ·  5 OF 5",
            "Can Aura override the 48-hour rule?",
            "No. Aura only sees that user’s own figures. The visa check still decides whether the shift is saved.",
            "Jenish: can Aura override the 48-hour rule?\nKrishna: does Aura write a shift row?\nSuraj: if Aura is off, do the numbers still show?",
            14,
        ),
    ]
    for item in questions:
        question_slide(prs, *item)
    for path in OUTS:
        path.parent.mkdir(parents=True, exist_ok=True)
        prs.save(path)
        print(f"Saved {path}")
    print(f"Slides: {len(prs.slides)}")


if __name__ == "__main__":
    build()
