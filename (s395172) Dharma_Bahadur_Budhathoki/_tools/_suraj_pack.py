"""Create Suraj folder: timesheet, data dictionary, timesheet figure."""
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Pt, Inches
from docx.oxml.ns import qn

ROOT = Path(r"C:\projects\clz\oraEarn\Suraj_Khadka_S394949")
ROOT.mkdir(parents=True, exist_ok=True)
FIGS = Path(r"C:\projects\clz\oraEarn\(s395172) Dharma_Bahadur_Budhathoki\figures")

ROWS = [
    ("Week 04", 1, "Read the OraEarn feature list from a numbers point of view: which screens would become measures (earnings, hours, expenses, savings) versus just labels."),
    ("Week 04", 1, "Sat with Krishna’s first entity list and marked which tables would feed a dashboard (shifts, expenses, savings goals) so analytics is not a separate made-up table."),
    ("Week 04", 1, "Drafted what net balance should mean: earnings in the period minus expenses in the same period, not a running bank balance."),
    ("Week 04", 1, "Noted the student visa fortnight as a separate measure (hours, not dollars) so it does not get mixed into the money cards."),
    ("Week 04", 1, "Listed questions for Jenish: which dashboard figures are in scope for the progress report, and which can wait."),
    ("Week 04", 1, "Group catch-up on roles. Confirmed I own dashboard meaning and Excel/PDF column sense, not the React screens themselves."),
    ("Week 05", 1, "Walked Krishna’s ERD and listed attributes that look like measures (hourly rate, total hours, total earned, amount, target, current amount)."),
    ("Week 05", 1, "Checked that export columns should come from 3NF tables (jobs, shifts, expenses) and not extra spreadsheet-only fields."),
    ("Week 05", 1, "Compared sample rows in Krishna’s notes with what a chart would need: date, job, hours, pay, expense category, amount."),
    ("Week 05", 1, "Flagged that SQL Server in the DBA notes and Postgres in the running app will confuse totals until one seed script is shared."),
    ("Week 05", 1, "Wrote expected empty-state vs seeded-state for the dashboard (no shifts = $0 earnings; Alice seed should not be an empty chart)."),
    ("Week 05", 1, "Jira: kept the analytics/export stories in view (dashboard EARN-106, later export EARN-147) so the role has a ticket, not only chat."),
    ("Week 06", 1, "Mapped fortnight hours to a rolling 14-day window, not a calendar fortnight, after reading the visa rule the app uses."),
    ("Week 06", 1, "Listed savings-rate as contributions versus target, and allocation % as a cap that should not go over 100% across goals."),
    ("Week 06", 1, "Reviewed expense kinds (fixed vs variable) for the spending chart so the dashboard is not one lump sum only."),
    ("Week 06", 1, "Asked for seed totals in writing so testers know if $1,440 expenses is Alice demo data."),
    ("Week 06", 1, "Looked at which joins a fortnight-hours query would need (user → shifts → work date) with Krishna."),
    ("Week 06", 1, "Short note on indexes that would help analytics reads (work date, user id) — support DBA, not a schema rewrite."),
    ("Week 07", 1, "Checked the live dashboard cards against the meaning list: net balance, earnings, expenses, savings targets, visa bar."),
    ("Week 07", 1, "Excel/PDF export (EARN-147): listed columns that must match the tables — date, job title, hours, pay, category, amount, payment."),
    ("Week 07", 1, "Noted 7-day vs 1-month vs 12-month filters change the money cards, so a screenshot must say which period."),
    ("Week 07", 1, "EARN-154 style test thinking: empty chart and seeded Alice chart are two different cases."),
    ("Week 07", 1, "Flagged that Git still has no folder under my name; timesheet and data dictionary need to exist as files."),
    ("Week 07", 1, "Read Dharma’s seeder notes (Cafe Barista $28, 14-day shifts) so expected hours are not guessed from the UI only."),
    ("Week 08", 1, "Re-checked visa 28/48 on the running dashboard against the 48-hour cap story."),
    ("Week 08", 2, "Went through export In Progress (EARN-147) and wrote which columns are in/out for week 8 so Jenish can name the right figure."),
    ("Week 08", 1, "Filled this timesheet and a one-page data dictionary so Assessment 2 has time-management evidence for the data-analyst seat."),
    ("Week 08", 1, "Agreed with the group that I still need a number-check on the live dashboard with Dharma before demo, and comments on the export ticket."),
    ("Week 08", 1, "Read Part E of the progress report in my name and confirmed the Jira keys and the honest Git gap."),
]

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Weeks 4-8"
headers = ["Week", "StudentId", "Name", "Hours [1-6]", "Activity details"]
head_font = Font(bold=True, color="FFFFFF")
head_fill = PatternFill("solid", fgColor="111111")
wrap = Alignment(wrap_text=True, vertical="center")
for i, h in enumerate(headers, 1):
    cell = ws.cell(1, i, h)
    cell.font = head_font
    cell.fill = head_fill
row_i = 2
last_week = None
total = 0
for week, hours, detail in ROWS:
    if last_week and week != last_week:
        row_i += 1
    ws.cell(row_i, 1, week)
    ws.cell(row_i, 2, "S394949")
    ws.cell(row_i, 3, "Suraj Khadka")
    ws.cell(row_i, 4, hours)
    c = ws.cell(row_i, 5, detail)
    c.alignment = wrap
    ws.row_dimensions[row_i].height = 36
    total += hours
    last_week = week
    row_i += 1
row_i += 1
ws.cell(row_i, 1, "Total").font = Font(bold=True)
ws.cell(row_i, 4, total).font = Font(bold=True)
ws.cell(row_i, 5, "Weeks 4–8. Cap 1–6 hours/week. Role: Data Analyst / supporting DBA.")
ws.column_dimensions["A"].width = 14
ws.column_dimensions["B"].width = 14
ws.column_dimensions["C"].width = 18
ws.column_dimensions["D"].width = 14
ws.column_dimensions["E"].width = 92
xlsx = ROOT / "Individual_timesheet_Suraj_Khadka_S394949_weeks4-8.xlsx"
wb.save(xlsx)
print("timesheet", xlsx, "hours", total)

# Data dictionary
doc = Document()
p = doc.add_paragraph()
run = p.add_run("OraEarn — one-page data dictionary (Suraj Khadka, S394949)")
run.bold = True
run.font.size = Pt(14)
doc.add_paragraph("For Assessment 2. These are the dashboard / export meanings I own checking. Not a full schema.")
rows = [
    ["Name on screen", "Comes from", "Meaning I am using"],
    ["Net balance", "Earnings − expenses in the selected period", "Not a bank account. Can be negative."],
    ["Earnings", "Sum of shift TotalEarned in the period", "Hours × job hourly rate after breaks."],
    ["Expenses", "Sum of expense Amount in the period", "By category (Rent, Food, Transport…)."],
    ["Savings targets / rate", "Goal current vs target, allocation %", "Allocation across goals should stay ≤ 100%."],
    ["Fortnight hours", "Sum of shift hours in rolling 14 days", "Student visa cap 48 hours. Separate from money."],
    ["Excel/PDF export", "Jobs, shifts, expenses tables", "Columns must exist on the 3NF tables."],
    ["Alice seed", "DemoDataSeeder (Cafe Barista $28, 14-day shifts)", "Charts should not be empty after seed."],
]
table = doc.add_table(rows=len(rows), cols=3)
table.style = "Table Grid"
for i, row in enumerate(rows):
    for j, v in enumerate(row):
        table.rows[i].cells[j].text = v
doc.add_paragraph("Jira: EARN-106 dashboard, EARN-147 Excel/PDF export (In Progress), EARN-154 analytics test cases.")
doc.add_paragraph("Git is still thin under my name. This file and the timesheet are the files I am adding for Assessment 2.")
dd = ROOT / "OraEarn_Data_Dictionary_Suraj_Khadka_S394949.docx"
doc.save(dd)
print("dictionary", dd)

# Timesheet figure
FONT_UI = ImageFont.truetype(r"C:\Windows\Fonts\segoeui.ttf", 16)
FONT_SM = ImageFont.truetype(r"C:\Windows\Fonts\segoeui.ttf", 12)
FONT_TH = ImageFont.truetype(r"C:\Windows\Fonts\segoeui.ttf", 13)
show = [("Week", "ID", "Name", "Hours", "Activity (short)")]
for week, hours, detail in ROWS:
    show.append((week, "S394949", "Suraj Khadka", str(hours), detail[:70] + "…"))
show = show[:22]
w, header_h, lh, pad = 1180, 72, 22, 16
img = Image.new("RGB", (w, header_h + pad + len(show) * lh + pad), (248, 248, 248))
d = ImageDraw.Draw(img)
d.rectangle((0, 0, w, header_h), fill=(0, 0, 0))
d.text((22, 10), "Timesheet proof — Suraj Khadka S394949 (weeks 4–8)", font=FONT_UI, fill=(255, 255, 255))
d.text((22, 34), "30 hours on file  ·  Data Analyst / supporting DBA  ·  1–6 hours/week cap", font=FONT_SM, fill=(160, 200, 220))
d.text((22, 52), "Folder Suraj_Khadka_S394949  ·  also a one-page data dictionary", font=FONT_SM, fill=(140, 140, 140))
y = header_h + 8
widths = [90, 90, 140, 60, 740]
for i, row in enumerate(show):
    bg = (0, 0, 0) if i == 0 else ((255, 255, 255) if i % 2 else (236, 236, 236))
    d.rectangle((12, y, w - 12, y + lh), fill=bg)
    x = 18
    fill = (255, 255, 255) if i == 0 else (30, 30, 30)
    for j, cell in enumerate(row):
        d.text((x, y + 3), cell, font=FONT_TH, fill=fill)
        x += widths[j]
    y += lh
fig = FIGS / "figure-timesheet-suraj.png"
img.save(fig)
print("figure", fig)
