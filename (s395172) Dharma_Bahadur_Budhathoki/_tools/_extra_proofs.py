"""Generate extra Assessment 2 proof images: git, timesheets, code glimpses."""
from __future__ import annotations

import subprocess
import textwrap
from pathlib import Path

import openpyxl
from PIL import Image, ImageDraw, ImageFont

FIGS = Path(r"C:\projects\clz\oraEarn\(s395172) Dharma_Bahadur_Budhathoki\figures")
REPO = Path(r"C:\projects\clz\oraEarn")
FONT_UI = ImageFont.truetype(r"C:\Windows\Fonts\segoeui.ttf", 16)
FONT_SM = ImageFont.truetype(r"C:\Windows\Fonts\segoeui.ttf", 12)
FONT_CODE = ImageFont.truetype(r"C:\Windows\Fonts\consola.ttf", 13)
FONT_TH = ImageFont.truetype(r"C:\Windows\Fonts\segoeui.ttf", 13)


def header(d: ImageDraw.ImageDraw, w: int, title: str, subtitle: str) -> None:
    d.rectangle((0, 0, w, 72), fill=(0, 0, 0))
    d.text((22, 10), title, font=FONT_UI, fill=(255, 255, 255))
    d.text((22, 34), subtitle, font=FONT_SM, fill=(160, 200, 220))
    d.text((22, 52), "Branch app-s395172  ·  Dharma001  ·  S395172  ·  PRT_681_SYDNEY_GROUP-4", font=FONT_SM, fill=(140, 140, 140))


def git_log_image() -> None:
    log = subprocess.check_output(
        ["git", "log", "-22", "--pretty=format:%h  %ad  %-14an  %s", "--date=short"],
        cwd=REPO,
        text=True,
        encoding="utf-8",
    )
    branches = subprocess.check_output(
        ["git", "branch", "-vv"],
        cwd=REPO,
        text=True,
        encoding="utf-8",
    )
    lines = ["Working branch vs main:", *branches.splitlines(), "", "Recent commits:", *log.splitlines()]
    w, pad, header_h, lh = 1100, 22, 72, 18
    img = Image.new("RGB", (w, header_h + pad + len(lines) * lh + pad), (18, 18, 18))
    d = ImageDraw.Draw(img)
    header(d, w, "Git proof — local clone of app-s395172", "Same history as GitHub PRT_681_SYDNEY_GROUP-4 (repo is private when logged out)")
    y = header_h + 8
    for line in lines:
        color = (120, 220, 140) if line.strip().startswith("*") else (225, 225, 225)
        d.text((pad, y), line[:120], font=FONT_CODE, fill=color)
        y += lh
    path = FIGS / "figure-git-log-detail.png"
    img.save(path)
    print("git", path.name, path.stat().st_size)


def timesheet_image(xlsx: Path, out: Path, title: str, subtitle: str, max_rows: int = 22) -> None:
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    rows: list[list[str]] = []
    for sheet in wb.sheetnames:
        ws = wb[sheet]
        for row in ws.iter_rows(max_row=min(ws.max_row or 1, 40), max_col=min(ws.max_column or 1, 6), values_only=True):
            cells = [("" if v is None else str(v).replace("\n", " "))[:90] for v in row]
            if any(c.strip() for c in cells):
                rows.append(cells)
        if len(rows) >= max_rows:
            break
    rows = rows[:max_rows]
    col_w = [170, 140, 220, 90, 420]
    while len(col_w) < max((len(r) for r in rows), default=1):
        col_w.append(160)
    w = min(sum(col_w[:5]) + 40, 1180)
    header_h, lh, pad = 72, 22, 16
    img = Image.new("RGB", (w, header_h + pad + (len(rows) + 1) * lh + pad), (248, 248, 248))
    d = ImageDraw.Draw(img)
    header(d, w, title, subtitle)
    y = header_h + 8
    for i, row in enumerate(rows):
        bg = (0, 0, 0) if i == 0 else ((255, 255, 255) if i % 2 else (236, 236, 236))
        d.rectangle((12, y, w - 12, y + lh), fill=bg)
        x = 18
        fill = (255, 255, 255) if i == 0 else (30, 30, 30)
        for j, cell in enumerate(row[:5]):
            d.text((x, y + 3), cell[:48], font=FONT_TH, fill=fill)
            x += col_w[j] if j < len(col_w) else 160
        y += lh
    img.save(out)
    print("timesheet", out.name, out.stat().st_size)


def glimpse(path: Path, title: str, subtitle: str, code: str) -> None:
    lines: list[str] = []
    for line in code.replace("\t", "    ").splitlines():
        if len(line) <= 100:
            lines.append(line)
        else:
            lines.extend(textwrap.wrap(line, 100, replace_whitespace=False) or [""])
    lines = lines[:18]
    w, pad, header_h, lh = 980, 22, 72, 18
    img = Image.new("RGB", (w, header_h + pad + len(lines) * lh + pad), (18, 18, 18))
    d = ImageDraw.Draw(img)
    header(d, w, title, subtitle)
    y = header_h + 8
    for line in lines:
        d.text((pad, y), line, font=FONT_CODE, fill=(225, 225, 225))
        y += lh
    img.save(path)
    print("glimpse", path.name)


def jira_keys_image() -> None:
    rows = [
        ("Key", "Title", "Column", "Owner area"),
        ("EARN-105", "Login Feature", "Done", "Dharma / auth"),
        ("EARN-106", "User Dashboard", "Done", "Dharma / Suraj analytics"),
        ("EARN-107", "Jobs CRUD", "Done", "Dharma"),
        ("EARN-70", "Log a shift manually / Shift CRUD", "Done", "Dharma (main feature)"),
        ("EARN-8", "Register with visa details and ID upload", "Idea", "Dharma / Jenish FR"),
        ("EARN-12", "Admin approve or reject pending registrations", "Idea", "Dharma"),
        ("EARN-147", "Export Excel and PDF", "In Progress", "Dharma / Suraj"),
        ("EARN-148", "Visa and analytics tests", "In Progress", "Krishna tester"),
        ("EARN-154", "Test cases visa analytics export", "In Progress", "Krishna / Suraj"),
        ("EARN-164", "Aura and notification tests", "To Do", "Dharma / tester"),
        ("EARN-165", "Intelligence UAT plan", "To Do", "Jenish"),
        ("EARN-144", "Week 7 timesheet meeting log", "Idea", "all members"),
        ("EARN-173", "Re-export ERD PNGs from draw.io", "Done", "Krishna"),
    ]
    w, header_h, lh, pad = 1080, 72, 24, 16
    img = Image.new("RGB", (w, header_h + pad + len(rows) * lh + pad), (248, 248, 248))
    d = ImageDraw.Draw(img)
    header(d, w, "Jira EARN keys used as evidence", "From ora--earn board screenshot — Stories only after Epic/Sub-task import failed")
    y = header_h + 8
    widths = [110, 420, 130, 360]
    for i, row in enumerate(rows):
        bg = (0, 0, 0) if i == 0 else ((255, 255, 255) if i % 2 else (236, 236, 236))
        d.rectangle((12, y, w - 12, y + lh), fill=bg)
        x = 20
        fill = (255, 255, 255) if i == 0 else (30, 30, 30)
        for j, cell in enumerate(row):
            d.text((x, y + 4), cell, font=FONT_TH, fill=fill)
            x += widths[j]
        y += lh
    path = FIGS / "figure-jira-keys.png"
    img.save(path)
    print("jira keys", path.name)


if __name__ == "__main__":
    FIGS.mkdir(exist_ok=True)
    git_log_image()
    jira_keys_image()
    timesheet_image(
        Path(r"C:\projects\clz\oraEarn\(s395172) Dharma_Bahadur_Budhathoki\Individual Timesheets - Dharma Bahadur Budhathoki(S395172) - weeks 4-8.xlsx"),
        FIGS / "figure-timesheet-dharma.png",
        "Timesheet proof — Dharma S395172 (weeks 4–8)",
        "29.5 hours on file  ·  1–6 hours/week cap",
        24,
    )
    timesheet_image(
        Path(r"C:\projects\clz\oraEarn\Krishna_sapkota_S395795\Individual_timesheet_Krishna_S395795_Timesheet.xlsx"),
        FIGS / "figure-timesheet-krishna.png",
        "Timesheet proof — Krishna S395795 (weeks 4–6 on file)",
        "18 hours on file  ·  weeks 7–8 still need the same sheet",
        20,
    )
    timesheet_image(
        Path(r"C:\projects\clz\oraEarn\Jenish_Gurung_s395321\IndividualTimesheet_JenishGurung_s395321.xlsx"),
        FIGS / "figure-timesheet-jenish.png",
        "Timesheet proof — Jenish s395321 (weeks 1–5 + week 8 extra)",
        "30 hours on the combined sheet  ·  week 8 extra 4 hours in a separate file",
        20,
    )
    glimpse(
        FIGS / "glimpse-seeder.png",
        "Backend glimpse — Demo seeder",
        "DemoDataSeeder.cs  ·  Cafe Barista $28, Retail $26, 14-day shifts",
        """private static async Task SeedAliceAsync(...)
{
    var cafe = new Job { Title = "Cafe Barista", HourlyRate = 28m };
    var retail = new Job { Title = "Retail Assistant", HourlyRate = 26m };
    db.Jobs.AddRange(cafe, retail);
    var fortnightStart = today.AddDays(-13);
    for (var offset = 0; offset < 14; offset++)
        AddShift(db, workDate, pattern.Job, pattern.Start, pattern.End, pattern.BreakMinutes);
}""",
    )
    glimpse(
        FIGS / "glimpse-shift-backend.png",
        "Backend glimpse — Shift CRUD",
        "ShiftService.cs  ·  create shift, calculate pay, visa 48-hour check",
        """public async Task<ShiftResponse> CreateAsync(string userId, ShiftRequest request)
{
    ValidateTimes(request.StartTime, request.EndTime);
    var (hours, earned) = Shift.Calculate(request.WorkDate, request.StartTime,
        request.EndTime, request.BreakMinutes, job.HourlyRate);
    await visaComplianceService.ValidateShiftAsync(userId, hours, request.WorkDate);
    db.Shifts.Add(new Shift { JobId = request.JobId, TotalHours = hours, TotalEarned = earned });
}""",
    )
    glimpse(
        FIGS / "glimpse-visa-backend.png",
        "Backend glimpse — 48-hour visa rule",
        "VisaComplianceService.cs  ·  student fortnight cap",
        """public const decimal StudentFortnightHourLimit = 48m;
public async Task ValidateShiftAsync(string userId, decimal shiftHours, DateOnly workDate)
{
    var hours = await SumFortnightHoursAsync(userId, fortnightStart, today);
    if (hours + shiftHours > StudentFortnightHourLimit)
        throw new InvalidOperationException("This shift would go over the 48-hour fortnight limit.");
}""",
    )
    print("OK")
