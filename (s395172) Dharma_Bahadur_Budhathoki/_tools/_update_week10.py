"""Add Week 9–10 to Dharma's timesheet and refresh the extract figure."""
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font
from PIL import Image, ImageDraw, ImageFont

DHARMA = Path(r"C:\projects\clz\oraEarn\(s395172) Dharma_Bahadur_Budhathoki")
OLD = DHARMA / "Individual Timesheets - Dharma Bahadur Budhathoki(S395172) - weeks 4-8.xlsx"
NEW = DHARMA / "Individual Timesheets - Dharma Bahadur Budhathoki(S395172) - weeks 4-10.xlsx"
FIG = DHARMA / "figures" / "figure-timesheet-dharma.png"

ROWS_W9 = [
    (1, "Packed the Assessment 2 Word file with labelled figures, GitHub proof, and Scrum Master diagrams."),
    (1, "Fixed Swagger register form so swagger.json returns 200, then ran live API tests (token hidden)."),
    (2, "Drew architecture, use-case, add-shift sequence and team ERD workshop pages after the group design talk."),
    (1, "Ran GET /health in Swagger (200) and Aura /docs tests. No demo passwords in the upload."),
    (1, "As Scrum Master, assembled the Learnline pack and kept the team on branch app-s395172."),
]
ROWS_W10 = [
    (1, "Updated my timesheet and renamed the Learnline submission files to Dharma Bahadur Budhathoki (S395172)."),
    (2, "Final review of the progress report, figures, and clickable GitHub / Jira links for Week 10 upload."),
    (1, "As Scrum Master, checked the private-repo invite note and that clone app-s395172 is still on the cover."),
    (1, "Prepared the Week 10 Learnline zip (report, my timesheet weeks 4-10, meeting logs)."),
    (1, "Polished remaining captions: Swagger tests, UML/ERD diagrams, and individual hours through Week 10."),
]


def font(size, bold=False):
    name = "C:/Windows/Fonts/calibrib.ttf" if bold else "C:/Windows/Fonts/calibri.ttf"
    try:
        return ImageFont.truetype(name, size)
    except OSError:
        return ImageFont.load_default()


def update_sheet():
    src = NEW if NEW.exists() else OLD
    wb = openpyxl.load_workbook(src)
    ws = wb.active
    # wipe old total row if present
    for r in range(ws.max_row, 1, -1):
        val = ws[f"A{r}"].value
        if val and str(val).startswith("Total"):
            ws.delete_rows(r, 1)
            break
    start = ws.max_row + 2
    r = start
    for hours, detail in ROWS_W9:
        ws[f"A{r}"] = "Week 09"
        ws[f"B{r}"] = "S395172"
        ws[f"C{r}"] = "Dharma Bahadur Budhathoki"
        ws[f"D{r}"] = hours
        ws[f"E{r}"] = detail
        r += 1
    r += 1
    for hours, detail in ROWS_W10:
        ws[f"A{r}"] = "Week 10"
        ws[f"B{r}"] = "S395172"
        ws[f"C{r}"] = "Dharma Bahadur Budhathoki"
        ws[f"D{r}"] = hours
        ws[f"E{r}"] = detail
        r += 1
    r += 1
    ws[f"A{r}"] = "Total hours (Week 04-10)"
    ws[f"B{r}"] = "S395172"
    ws[f"C{r}"] = "Dharma Bahadur Budhathoki"
    ws[f"D{r}"] = "=SUM(D2:D{0})".format(r - 1)
    ws[f"E{r}"] = "Must stay within 1-6 hours per week. Week 04=5.5, Weeks 05-10=6 each. Total 41.5."
    for cell in ws[f"A{r}:E{r}"]:
        for c in cell:
            c.font = Font(bold=True)
    ws.column_dimensions["E"].width = 110
    for row in ws.iter_rows(min_row=2, max_col=5):
        row[4].alignment = Alignment(wrap_text=True)
    wb.save(NEW)
    print("saved", NEW.name, "total row", r)
    return NEW


def draw_figure(path: Path):
    wb = openpyxl.load_workbook(path, data_only=False)
    ws = wb.active
    rows = []
    for row in ws.iter_rows(min_row=1, max_col=5, values_only=True):
        if any(row):
            week, sid, name, hours, detail = row
            detail = (str(detail)[:72] + "…") if detail and len(str(detail)) > 72 else (detail or "")
            rows.append((str(week or ""), str(sid or ""), str(name or ""), str(hours or ""), str(detail)))
    W, row_h = 1600, 28
    H = 70 + row_h * len(rows) + 20
    im = Image.new("RGB", (W, H), (255, 255, 255))
    d = ImageDraw.Draw(im)
    fb, f = font(22, True), font(14)
    d.text((20, 12), "Timesheet — Dharma Bahadur Budhathoki (S395172)  weeks 4–10", font=fb, fill=(28, 28, 28))
    d.text((20, 42), "41.5 hours on file. Cap 1–6 hours/week. Branch app-s395172.", font=f, fill=(90, 90, 90))
    cols = [20, 140, 250, 520, 610]
    y = 70
    for i, row in enumerate(rows):
        if i == 0 or str(row[0]).startswith("Total"):
            d.rectangle((16, y, W - 16, y + row_h), fill=(230, 230, 230))
        elif i % 2 == 0:
            d.rectangle((16, y, W - 16, y + row_h), fill=(245, 245, 245))
        for x, cell in zip(cols, row):
            d.text((x, y + 6), cell, font=font(13, i == 0), fill=(28, 28, 28))
        y += row_h
    im.save(FIG)
    print("wrote", FIG.name, FIG.stat().st_size)


if __name__ == "__main__":
    p = update_sheet()
    draw_figure(p)
