"""Build the Assessment 4 handover as a Word document for Dharma."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT = Path(
    r"C:\projects\clz\oraEarn\oraEarn-WebApp\docs"
    r"\PRT681_Assessment4_Dharma_Bahadur_Budhathoki_S395172_OraEarn_Handover.docx"
)
JIRA = (
    "https://students-team-dh8zehye.atlassian.net/jira/core/projects/EARN/board"
    "?filter=&groupBy=none"
)
GIT = "https://github.com/itsmekrishnaaaaaa/PRT_681_SYDNEY_GROUP-4"
FIG = Path(r"C:\projects\clz\oraEarn\(s395172) Dharma_Bahadur_Budhathoki\figures")
LOGO_SRC = Path(
    r"C:\Users\budha\.cursor\projects\c-projects-clz-oraEarn-oraEarn-WebApp\assets"
    r"\c__Users_budha_AppData_Roaming_Cursor_User_workspaceStorage_e1912c2036a7fec9ab349d4dfa8105e7_images_image-06f37480-90d4-4dc0-b096-c74bb047362e.png"
)
LOGO = FIG / "cdu-logo.png"
NAME = "Dharma Bahadur Budhathoki"
SID = "S395172"
GROUP = "Group 4"

NAVY = RGBColor(0x1F, 0x3A, 0x5F)
MUTED = RGBColor(0x55, 0x55, 0x55)
LINK = "1F4E79"
_bookmark_id = 1


def _shade(cell, fill: str) -> None:
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def _set_run_font(run, name: str, size: int, bold: bool = False, color=None, italic: bool = False) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def add_hyperlink(paragraph, text: str, url: str) -> None:
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), LINK)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rFonts = OxmlElement("w:rFonts")
    rFonts.set(qn("w:ascii"), "Calibri")
    rFonts.set(qn("w:hAnsi"), "Calibri")
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "22")
    rPr.extend([rFonts, color, underline, sz])
    run.append(rPr)
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    run.append(t)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_bookmark(paragraph, name: str) -> None:
    global _bookmark_id
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(_bookmark_id))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(_bookmark_id))
    paragraph._p.insert(0, start)
    paragraph._p.append(end)
    _bookmark_id += 1


def add_internal_link(paragraph, text: str, anchor: str) -> None:
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("w:anchor"), anchor)
    run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), LINK)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rFonts = OxmlElement("w:rFonts")
    rFonts.set(qn("w:ascii"), "Calibri")
    rFonts.set(qn("w:hAnsi"), "Calibri")
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "22")
    rPr.extend([rFonts, color, underline, sz])
    run.append(rPr)
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    run.append(t)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def style_body(paragraph) -> None:
    pf = paragraph.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(8)
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = 1.08


def add_p(doc: Document, text: str):
    p = doc.add_paragraph()
    style_body(p)
    run = p.add_run(text)
    _set_run_font(run, "Calibri", 11)
    return p


def add_mixed(doc: Document, parts: list[tuple[str, str]]):
    """parts: ('text', str) or ('link', text, url) or ('anchor', text, bookmark)."""
    p = doc.add_paragraph()
    style_body(p)
    for part in parts:
        if part[0] == "text":
            run = p.add_run(part[1])
            _set_run_font(run, "Calibri", 11)
        elif part[0] == "link":
            add_hyperlink(p, part[1], part[2])
        elif part[0] == "anchor":
            add_internal_link(p, part[1], part[2])
    return p


def add_h(doc: Document, text: str, level: int, bookmark: str):
    p = doc.add_heading(text, level=level)
    add_bookmark(p, bookmark)
    for run in p.runs:
        run.font.color.rgb = NAVY
        run.font.name = "Calibri"
    p.paragraph_format.space_before = Pt(14 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(6)
    return p


def add_author(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    style_body(p)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    _set_run_font(run, "Calibri", 10, italic=True, color=MUTED)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(item, style="List Bullet")
        style_body(p)
        for run in p.runs:
            _set_run_font(run, "Calibri", 11)


def add_code(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.3)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(text)
    _set_run_font(run, "Consolas", 9)
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), "F4F6F8")
    pPr.append(shd)


def add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        run = p.add_run(header)
        _set_run_font(run, "Calibri", 10, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
        _shade(cell, "1F3A5F")
    for r, row in enumerate(rows, start=1):
        for c, value in enumerate(row):
            cell = table.rows[r].cells[c]
            cell.text = ""
            p = cell.paragraphs[0]
            run = p.add_run(value)
            _set_run_font(run, "Calibri", 10)
            if r % 2 == 0:
                _shade(cell, "F4F6F8")
    doc.add_paragraph().paragraph_format.space_after = Pt(6)


def prepare_logo() -> None:
    if not LOGO_SRC.exists():
        return
    try:
        from PIL import Image
    except ImportError:
        return
    image = Image.open(LOGO_SRC).convert("RGBA")
    pixels = []
    for red, green, blue, alpha in image.getdata():
        if red < 28 and green < 28 and blue < 28:
            pixels.append((red, green, blue, 0))
        else:
            pixels.append((red, green, blue, alpha))
    image.putdata(pixels)
    bbox = image.getbbox()
    if bbox:
        image = image.crop(bbox)
    LOGO.parent.mkdir(parents=True, exist_ok=True)
    image.save(LOGO)


def add_page_break(doc: Document) -> None:
    doc.add_page_break()


def add_caption(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    style_body(p)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    _set_run_font(run, "Calibri", 9, italic=True, color=MUTED)


FIGURES = [
    ("figure-jira-earn-board.png", "EARN board, 16 October 2026. The work is in Done. DB on a card is me, Dharma Bahadur Budhathoki (S395172)."),
    ("figure-diagram-architecture.png", "How the browser, the API, the database, and Aura sit next to each other."),
    ("figure-folder-structure.png", "The project folders. Each feature has its own place in the API and in the React app."),
    ("figure-login-desktop.png", "Login. The eye shows the password. The database stores the hash, not the typed password."),
    ("figure-register-desktop.png", "Registration. The account stays Pending until an admin approves it."),
    ("figure-github-repo-live.png", "Public repository PRT_681_SYDNEY_GROUP-4. I am Dharma001 in the contributor list."),
    ("figure-git-log-detail.png", "Git history for the work on this branch. One commit, one reason."),
    ("figure-backend-db.png", "The database side of the API. Shifts point at a job. The job points at the user."),
    ("figure-swagger-api.png", "Swagger on port 5203. I try an endpoint here before I wire the React screen."),
    ("figure-swagger-health.png", "GET /health returns 200 when the API is up."),
    ("glimpse-visa-backend.png", "The 48-hour check runs in the API before a student shift is saved."),
    ("figure-diagram-swagger-tests.png", "The tests I run before I call a feature finished."),
    ("figure-diagram-team-erd.png", "The tables. A shift points at a job. The job points at the user."),
    ("figure-diagram-uml-classes.png", "Class diagram for the same model."),
    ("figure-jobs-desktop.png", "Jobs. The hourly rate on the job is what a shift gets paid."),
    ("figure-jobs-add-desktop.png", "Add-job dialog. The message under the field is the browser check."),
    ("figure-jobs-mobile.png", "The same jobs page on a narrow screen. The layout stacks. It does not hide the form."),
    ("figure-shifts-desktop.png", "Shift calendar. Hours and pay are calculated on the server. The visa bar is on this page."),
    ("figure-dashboard-desktop.png", "Dashboard. Net, earnings, and expenses for the period I pick."),
    ("figure-expenses-desktop.png", "Expenses, with category and type."),
    ("figure-savings-desktop.png", "Savings goals. A goal takes a share of leftover pay."),
    ("figure-timesheet-dharma.png", "My timesheet, Dharma Bahadur Budhathoki, S395172."),
]


def add_figure(doc: Document, filename: str) -> None:
    comment = next(text for name, text in FIGURES if name == filename)
    number = next(i for i, (name, _) in enumerate(FIGURES, start=1) if name == filename)
    path = FIG / filename
    if not path.exists():
        return
    pic = doc.add_paragraph()
    pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic.paragraph_format.space_before = Pt(6)
    pic.paragraph_format.space_after = Pt(2)
    run = pic.add_run()
    run.add_picture(str(path), width=Cm(14.6))

    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(10)
    add_bookmark(cap, f"fig{number}")
    text = cap.add_run(f"Figure {number}. {comment}")
    _set_run_font(text, "Calibri", 9, italic=True, color=MUTED)


def insert_toc(paragraph) -> None:
    run = paragraph.add_run()
    fld = OxmlElement("w:fldChar")
    fld.set(qn("w:fldCharType"), "begin")
    run._r.append(fld)

    run = paragraph.add_run()
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = ' TOC \\o "1-2" \\h \\z \\u '
    run._r.append(instr)

    run = paragraph.add_run()
    fld = OxmlElement("w:fldChar")
    fld.set(qn("w:fldCharType"), "separate")
    run._r.append(fld)

    run = paragraph.add_run("Contents update when the file is opened in Word.")
    _set_run_font(run, "Calibri", 11, italic=True, color=MUTED)

    run = paragraph.add_run()
    fld = OxmlElement("w:fldChar")
    fld.set(qn("w:fldCharType"), "end")
    run._r.append(fld)


def insert_tof(doc: Document) -> None:
    for number, (_, comment) in enumerate(FIGURES, start=1):
        p = doc.add_paragraph()
        style_body(p)
        p.paragraph_format.space_after = Pt(2)
        add_internal_link(p, f"Figure {number}. {comment}", f"fig{number}")
        gap = p.add_run("    ")
        _set_run_font(gap, "Calibri", 11)
        add_page_field(p, f"PAGEREF fig{number}")


def add_page_field(paragraph, instr_text: str) -> None:
    run = paragraph.add_run()
    fld = OxmlElement("w:fldChar")
    fld.set(qn("w:fldCharType"), "begin")
    run._r.append(fld)
    run = paragraph.add_run()
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" {instr_text} "
    run._r.append(instr)
    run = paragraph.add_run()
    fld = OxmlElement("w:fldChar")
    fld.set(qn("w:fldCharType"), "end")
    run._r.append(fld)


def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    for style_name, size in (("Heading 1", 16), ("Heading 2", 13)):
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = NAVY
        style.font.bold = True


def set_header_footer(section, header_text: str, with_page: bool) -> None:
    section.different_first_page_header_footer = False
    header = section.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = hp.add_run(header_text)
    _set_run_font(run, "Calibri", 9, color=MUTED)

    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if with_page:
        left = fp.add_run("Group 4  |  PRT_681_SYDNEY_GROUP-4    Page ")
        _set_run_font(left, "Calibri", 9, color=MUTED)
        add_page_field(fp, "PAGE")
        mid = fp.add_run(" of ")
        _set_run_font(mid, "Calibri", 9, color=MUTED)
        add_page_field(fp, "NUMPAGES")
    else:
        run = fp.add_run("Group 4  |  PRT_681_SYDNEY_GROUP-4  |  Charles Darwin University")
        _set_run_font(run, "Calibri", 9, color=MUTED)

    # footer page-number runs were added without font; set them
    for run in fp.runs:
        if run.font.size is None:
            _set_run_font(run, "Calibri", 9, color=MUTED)


def enable_update_fields(doc: Document) -> None:
    settings = doc.settings.element
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    settings.append(update)


def build() -> None:
    doc = Document()
    configure_styles(doc)
    for section in doc.sections:
        section.top_margin = Cm(1.8)
        section.bottom_margin = Cm(1.8)
        section.left_margin = Cm(2.0)
        section.right_margin = Cm(2.0)
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)

    # Cover uses the first section. No page number on the cover.
    cover = doc.sections[0]
    set_header_footer(cover, "", with_page=False)
    cover.header.paragraphs[0].text = ""
    cover.footer.is_linked_to_previous = False

    prepare_logo()
    if LOGO.exists():
        logo = doc.add_paragraph()
        logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        logo.paragraph_format.space_before = Pt(36)
        logo.paragraph_format.space_after = Pt(6)
        logo.add_run().add_picture(str(LOGO), width=Cm(7.2))

    uni = doc.add_paragraph()
    uni.alignment = WD_ALIGN_PARAGRAPH.CENTER
    uni.paragraph_format.space_before = Pt(10)
    uni.paragraph_format.space_after = Pt(0)
    run = uni.add_run("CHARLES DARWIN UNIVERSITY")
    _set_run_font(run, "Calibri", 11, bold=True, color=NAVY)

    faculty = doc.add_paragraph()
    faculty.alignment = WD_ALIGN_PARAGRAPH.CENTER
    faculty.paragraph_format.space_after = Pt(14)
    run = faculty.add_run("Faculty of Science and Technology")
    _set_run_font(run, "Calibri", 11, color=MUTED)

    rule = doc.add_paragraph()
    rule.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rule.paragraph_format.space_before = Pt(4)
    rule.paragraph_format.space_after = Pt(8)
    run = rule.add_run("PRT681  |  Software Project  |  Assessment 4")
    _set_run_font(run, "Calibri", 12, bold=True, color=NAVY)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(16)
    title.paragraph_format.space_after = Pt(2)
    run = title.add_run("OraEarn")
    _set_run_font(run, "Calibri", 36, bold=True, color=NAVY)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.paragraph_format.space_after = Pt(14)
    run = sub.add_run("Handover document")
    _set_run_font(run, "Calibri", 16, color=MUTED)

    add_table(
        doc,
        ["Detail", "Entry"],
        [
            ["Student", "Dharma Bahadur Budhathoki"],
            ["Student number", "S395172"],
            ["Role", "Developer"],
            ["Group", "Group 4, Sydney"],
            ["Unit", "PRT681 Software Project"],
            ["Date", "16 October 2026"],
        ],
    )

    git_line = doc.add_paragraph()
    git_line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_body(git_line)
    run = git_line.add_run("Repository  ")
    _set_run_font(run, "Calibri", 11)
    add_hyperlink(git_line, "PRT_681_SYDNEY_GROUP-4", GIT)

    jira_line = doc.add_paragraph()
    jira_line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_body(jira_line)
    run = jira_line.add_run("Jira board  ")
    _set_run_font(run, "Calibri", 11)
    add_hyperlink(jira_line, "EARN project", JIRA)

    # Body section so the cover is not numbered with the report pages.
    body = doc.add_section()
    body.top_margin = Cm(1.8)
    body.bottom_margin = Cm(1.8)
    body.left_margin = Cm(2.0)
    body.right_margin = Cm(2.0)
    set_header_footer(body, "Group 4  |  PRT681 Assessment 4  |  OraEarn", with_page=True)
    pg = OxmlElement("w:pgNumType")
    pg.set(qn("w:start"), "1")
    body._sectPr.append(pg)

    add_h(doc, "Contents", 1, "contents")
    toc_p = doc.add_paragraph()
    insert_toc(toc_p)
    add_page_break(doc)

    add_h(doc, "Table of figures", 1, "figures")
    insert_tof(doc)
    add_page_break(doc)

    add_h(doc, "1. What the submission asks for", 1, "what-submission-asks")
    add_author(doc, "Written by Dharma Bahadur Budhathoki (S395172)")
    add_p(
        doc,
        "Assessment 4 asks us to hand the project over. The marker wants the working app, and a write-up "
        "someone else can follow without ringing me. Three parts are required, and each part has to name who wrote it.",
    )
    add_bullets(
        doc,
        [
            "URLs. The git repository, the Jira board, the learning documents, and each person's timesheet.",
            "Set-up instructions. Tools, libraries, architecture, sign-in, git, the database, Entity Framework, connection strings, first run, dependency injection, exceptions, tests, and the data model.",
            "Adding a new feature. From the Jira card, through the code, a new service, libraries, styles, validation, a test, and the commit.",
        ],
    )
    add_p(
        doc,
        "Where I have put a screenshot, the caption under it is the comment. The list of those pictures is in the table of figures.",
    )

    add_h(doc, "2. URLs", 1, "urls")
    add_author(doc, "Written by Dharma Bahadur Budhathoki (S395172). The requirements files are Jenish's.")
    add_p(doc, "These are the links I actually use. The Jira one opens the EARN board, which is where my developer cards live.")
    add_table(
        doc,
        ["What", "Where"],
        [
            ["Jira board (EARN)", "students-team-dh8zehye.atlassian.net — EARN board"],
            ["Git repository", "github.com/itsmekrishnaaaaaa/PRT_681_SYDNEY_GROUP-4"],
            ["Requirements", "oraEarn-WebApp/docs/OraEarn-Requirements.md"],
            ["Project proposal", "oraEarn-WebApp/docs/OraEarn-Project-Proposal.md"],
            ["Normalisation (1NF to 3NF)", "oraEarn-WebApp/docs/OraEarn-Normalization-1NF-to-3NF.md"],
            ["ERD", "oraEarn-WebApp/docs/OraEarn-ERD-3NF.drawio"],
            ["My timesheet, weeks 4 to 12", "Dharma Bahadur Budhathoki (S395172), weeks 4-12.xlsx"],
            ["Jenish timesheet", "Jenish_Gurung_s395321/IndividualTimesheet_JenishGurung_s395321.xlsx"],
            ["Krishna timesheet", "Krishna_sapkota_S395795/Individual_timesheet_Krishna_S395795.xlsx"],
            ["Group timesheet", "Group4_Timesheet_PRT681.xlsx"],
        ],
    )
    add_mixed(
        doc,
        [
            ("text", "Jira: "),
            ("link", JIRA, JIRA),
        ],
    )
    add_mixed(
        doc,
        [
            ("text", "Git: "),
            ("link", GIT, GIT),
        ],
    )
    add_p(doc, "When I am running it on my own machine, these are the addresses.")
    add_table(
        doc,
        ["Service", "Address"],
        [
            ["React app", "http://localhost:5173"],
            ["API and Swagger", "http://localhost:5203/swagger"],
            ["Health", "http://localhost:5203/health"],
            ["SignalR hub", "http://localhost:5203/hubs/shifts"],
            ["Aura, optional", "http://localhost:8000"],
        ],
    )
    add_p(
        doc,
        "Once the API has seeded the demo data, I sign in as alice@oraearn.local with User123! when I want the student. "
        "Bob uses the same password and is a citizen, so the 48-hour rule does not apply to him. "
        "The admin is admin@oraearn.local with Admin123!.",
    )
    add_p(
        doc,
        "The README at the top of the GitHub repository is the short version of this handover. "
        "It carries the same GitHub link, the same Jira link, the three run commands, and the demo sign-ins. "
        "My name on that page is Dharma Bahadur Budhathoki, student number S395172.",
    )
    add_figure(doc, "figure-jira-earn-board.png")

    add_h(doc, "3. Set-up instructions", 1, "setup")
    add_p(
        doc,
        "I wrote this for a clean Windows laptop and PowerShell. The bits I touch every day are mine. "
        "Krishna's database steps and Suraj's data model are here too, so the next person is not stuck waiting on us.",
    )

    add_h(doc, "3.1 Tools and versions", 2, "tools")
    add_author(doc, "Written by Dharma Bahadur Budhathoki (S395172)")
    add_p(doc, "Install these first. If dotnet, node, or npm is missing, setup.ps1 stops and tells you.")
    add_table(
        doc,
        ["Tool", "Version", "Used for"],
        [
            [".NET SDK", "10 (net10.0)", "API, Identity, Entity Framework"],
            ["Node.js", "20 or newer", "React app"],
            ["PostgreSQL", "16 or newer, port 5432", "The only database"],
            ["Python", "3.11 or newer", "Aura. The rest of the app runs without it"],
            ["Git", "current", "The repository above"],
        ],
    )
    add_code(doc, "dotnet --version\nnode --version\npython --version")

    add_h(doc, "3.2 Libraries and versions", 2, "libraries")
    add_author(doc, "Written by Dharma Bahadur Budhathoki (S395172)")
    add_p(doc, "I copied these out of the project files so the versions are the ones we actually build with. Add a package with npm or dotnet add. Don't edit the lock file by hand.")
    add_table(
        doc,
        ["API package", "Version"],
        [
            ["Microsoft.AspNetCore.Authentication.JwtBearer", "10.0.11"],
            ["Microsoft.AspNetCore.Identity.EntityFrameworkCore", "10.0.11"],
            ["Microsoft.EntityFrameworkCore.Design", "10.0.11"],
            ["Npgsql.EntityFrameworkCore.PostgreSQL", "10.0.3"],
            ["Swashbuckle.AspNetCore", "8.1.4"],
        ],
    )
    add_table(
        doc,
        ["Frontend package", "Version"],
        [
            ["react, react-dom", "19.2"],
            ["vite", "8.2"],
            ["typescript", "6.0"],
            ["react-router-dom", "7.18"],
            ["axios", "1.19"],
            ["@microsoft/signalr", "10.0"],
            ["tailwindcss", "4.3"],
        ],
    )
    add_p(doc, "Tests use xUnit 2.9.3. Aura uses FastAPI 0.115 or newer, Uvicorn 0.32 or newer, and Pydantic 2.10 or newer.")

    add_h(doc, "3.3 Architecture", 2, "architecture")
    add_author(doc, "Written by Dharma Bahadur Budhathoki (S395172)")
    add_p(
        doc,
        "There are three programs. I did not glue them into one. The website only calls the API. "
        "The API talks to PostgreSQL, and it calls Aura if that Python process is running. "
        "If Aura is off, shifts, money, and the visa bar still work. The chat just says it could not get an answer.",
    )
    add_code(
        doc,
        "Browser (React, port 5173)\n"
        "    ->  ASP.NET Core API (port 5203)   JWT, REST, SignalR\n"
        "            ->  PostgreSQL\n"
        "            ->  Aura, Python FastAPI (port 8000)   optional",
    )
    add_p(
        doc,
        "On the API, each feature has its own folder: Auth, Jobs, Shifts, Expenses, SavingsGoals, "
        "Analytics, Notifications, Admin, VisaCompliance, and AiAdvisor. The class, the DTOs, the service, "
        "and the controller sit together. Program.cs stays short. I register everything in ServiceCollectionExtensions.",
    )
    add_figure(doc, "figure-diagram-architecture.png")
    add_figure(doc, "figure-folder-structure.png")
    add_p(
        doc,
        "On the website, each screen is under src/features. Sign-in is under src/auth. "
        "The axios client is src/shared/api/client.ts. The form checks are under src/shared/form-field. "
        "I put the routes in App.tsx.",
    )

    add_h(doc, "3.4 Authentication and authorisation", 2, "auth")
    add_author(doc, "Written by Dharma Bahadur Budhathoki (S395172)")
    add_p(
        doc,
        "Sign-in is who you are. The role is what you are allowed to do. I built both. Hiding a page in React is not the same as locking the API.",
    )
    add_bullets(
        doc,
        [
            "Register is POST /api/auth/register. The form includes the visa file. I save the account as Pending. I don't give out a token until an admin approves it.",
            "Login is POST /api/auth/login. A good login returns an access token that lasts 15 minutes and a refresh token that lasts 7 days.",
            "The website stores both in localStorage under the key auth, and sends Authorization: Bearer on each call.",
            "If a call comes back 401, the client tries the refresh token once. Logout wipes the refresh token on the user row.",
            "Passwords go through ASP.NET Identity. I ask for 8 characters, an upper letter, a lower letter, and a digit. Five wrong tries lock the account for 15 minutes. The table stores PasswordHash. It never stores the password the person typed.",
            "The token must say issuer OraEarn and audience OraEarnClient, and the secret has to be at least 32 characters. SignalR reads the same token from the access_token query on /hubs, because the socket can't set the header.",
        ],
    )
    add_table(
        doc,
        ["Place", "Rule"],
        [
            ["Almost every controller", "[Authorize]. No token is 401."],
            ["AdminController", "[Authorize(Roles = Admin)]. A student token is 403."],
            ["ProtectedRoute", "A signed-out person goes to /login. An admin goes to /admin."],
            ["AdminRoute", "A student goes back to /dashboard."],
            ["Services", "Queries use the user id inside the token. One student cannot read another student's pay."],
        ],
    )
    add_p(doc, "If I add a new endpoint, it still needs [Authorize]. Swagger will happily call anything the API leaves open.")
    add_figure(doc, "figure-login-desktop.png")
    add_figure(doc, "figure-register-desktop.png")

    add_h(doc, "3.5 Git download", 2, "git")
    add_author(doc, "Written by Dharma Bahadur Budhathoki (S395172)")
    add_p(doc, "Clone it once, then make a branch. I don't commit straight to main.")
    add_code(doc, "git clone " + GIT + "\ncd PRT_681_SYDNEY_GROUP-4\ngit checkout -b OE-<ticket>-short-name")
    add_p(
        doc,
        "The git root is the folder above oraEarn-WebApp. Open oraEarn-WebApp in the editor. "
        "Don't commit appsettings.Local.json, frontend/.env, node_modules, bin, obj, or the Python .venv. "
        "On the board I name the card with an EARN key.",
    )
    add_figure(doc, "figure-github-repo-live.png")
    add_figure(doc, "figure-git-log-detail.png")

    add_h(doc, "3.6 Database setup", 2, "database")
    add_author(doc, "Written by Krishna Sapkota (S395795). I checked it against the API on my machine. Dharma Bahadur Budhathoki, S395172.")
    add_p(
        doc,
        "PostgreSQL has to be running on port 5432. The first time the API starts, it connects to the postgres database, "
        "creates oraearn if it is not there yet, runs the migrations, and in Development it seeds the roles, the payment methods, and the demo users.",
    )
    add_code(doc, 'Get-Service postgresql* | Start-Service')
    add_p(
        doc,
        "On my machine the user is postgres and the password is 12345. If yours is different, change only "
        "backend/OraEarn.Api/appsettings.Local.json. After the first start you should see Jobs, Shifts, "
        "ActiveShiftSessions, ExpenseCategories, PaymentMethods, Expenses, SavingsGoals, and UserNotificationReads, plus the Identity tables. "
        "Shifts does not have a UserId column. You get the worker through the job.",
    )
    add_figure(doc, "figure-backend-db.png")

    add_h(doc, "3.7 Entity Framework commands", 2, "ef")
    add_author(doc, "Written by Krishna Sapkota (S395795). The migration names are the ones in the repo.")
    add_p(doc, "You don't need to run a migration just to start the app. Startup already calls MigrateAsync. After you change a model, from backend/OraEarn.Api:")
    add_code(
        doc,
        "dotnet tool install --global dotnet-ef\n"
        "dotnet ef migrations add <Name> --project OraEarn.Api.csproj\n"
        "dotnet ef database update --project OraEarn.Api.csproj",
    )
    add_p(
        doc,
        "Name the migration for the reason, like AddLeaveRequests. Read the Up method before you commit it. "
        "Deleting a user also deletes that person's jobs, categories, goals, and the live clock-in. "
        "You can't delete a job while shifts still point at it, or a category while expenses still point at it. "
        "The tests use an in-memory database, so dotnet test does not need PostgreSQL.",
    )

    add_h(doc, "3.8 Connection strings", 2, "connection")
    add_author(doc, "Written by Krishna Sapkota (S395795). On my machine I only edit the Local file. Dharma Bahadur Budhathoki, S395172.")
    add_table(
        doc,
        ["File", "Committed", "What belongs here"],
        [
            ["appsettings.json", "Yes", "Empty secret, empty CORS, demo seed off"],
            ["appsettings.Development.json", "Yes", "CORS for port 5173, development secret, demo seed on"],
            ["appsettings.Local.json", "No", "Your Postgres host, user, and password"],
            ["Environment variables", "On the server", "Production connection string, JWT, CORS, seed"],
        ],
    )
    add_code(
        doc,
        '{\n  "ConnectionStrings": {\n'
        '    "DefaultConnection": "Host=localhost;Port=5432;Database=oraearn;Username=postgres;Password=12345"\n'
        "  }\n}",
    )
    add_p(
        doc,
        "If DefaultConnection is missing, the API stops and says so. On a real server, set ConnectionStrings__DefaultConnection, "
        "Jwt__Secret (at least 32 characters, and not the development one), Cors__AllowedOrigins__0, Seed__DemoData=false, "
        "and the admin email and password. Don't put Aura on the public internet.",
    )

    add_h(doc, "3.9 How to run the applications", 2, "run")
    add_author(doc, "Written by Dharma Bahadur Budhathoki (S395172)")
    add_p(doc, "From the oraEarn-WebApp folder, run this once:")
    add_code(doc, ".\\scripts\\setup.ps1")
    add_p(doc, "That copies frontend/.env and appsettings.Local.json if they are missing, installs the npm packages, restores the API, and makes the Aura virtual environment if Python is there. Then I open three terminals:")
    add_code(doc, "cd backend\n.\\run.ps1\n\ncd frontend\n.\\run.ps1\n\ncd ai-advisor\n.\\run.ps1")
    add_p(
        doc,
        "Open http://localhost:5173 and sign in as Alice. Start the API before you open the browser. "
        "Aura is the third terminal. You can skip it and the rest of the app still runs. "
        "frontend/.env only needs VITE_API_URL=http://localhost:5203/api and VITE_HUB_URL=http://localhost:5203/hubs/shifts.",
    )
    add_figure(doc, "figure-swagger-api.png")
    add_figure(doc, "figure-swagger-health.png")

    add_h(doc, "3.10 Dependency injection", 2, "di")
    add_author(doc, "Written by Dharma Bahadur Budhathoki (S395172)")
    add_p(
        doc,
        "Program.cs calls AddApplicationServices. That is the only place I register a service, and each one lives for a single request. "
        "The list is AuthService, TokenService, JobService, ShiftService, ActiveShiftSessionService, ShiftNotificationService, "
        "ExpenseCategoryService, ExpenseService, SavingsGoalService, SavingsGoalSyncService, SavingsGoalAnalyticsService, "
        "SavingsAdvisorService, AdminService, AnalyticsService, NotificationService, VisaComplianceService, "
        "RegistrationApprovalAdvisorService, RegistrationApprovalService, and VisaDocumentStorageService.",
    )
    add_p(
        doc,
        "ApplicationDbContext is registered with AddDbContext. PythonAdvisorClient uses AddHttpClient, so it gets a proper HttpClient instead of one I create myself. "
        "A controller takes the service in its constructor. I don't write new JobService() inside a controller. "
        "If I forget the AddScoped line, the API can't find the service and it won't start that request.",
    )

    add_h(doc, "3.11 Exception handling", 2, "exceptions")
    add_author(doc, "Written by Dharma Bahadur Budhathoki (S395172)")
    add_p(doc, "When something the user did is wrong, I catch it in the controller and send a normal HTTP response. I don't let it fall through as a raw 500.")
    add_table(
        doc,
        ["What happened", "What the code does", "HTTP"],
        [
            ["Empty title, rate of zero, bad allocation", "ArgumentException, body { message }", "400"],
            ["Student shift would pass 48 hours", "InvalidOperationException from the visa check, before the insert", "400"],
            ["Row missing, or not this user's row", "Service returns null", "404"],
            ["No user id on the token", "UnauthorizedAccessException", "401"],
            ["Aura is down", "Client logs a warning and returns null. The API still answers.", "200"],
            ["Postgres is down at startup", "Message points at appsettings.Local.json", "Process stops"],
            ["JWT secret missing or still the dev value outside Development", "Thrown while building services", "Process stops"],
        ],
    )
    add_p(
        doc,
        "On the website the dialog stays open. The job dialog copies the API message into the summary at the top. "
        "A failed save does not wipe what the person typed. In production, anything I didn't catch goes through UseExceptionHandler.",
    )
    add_figure(doc, "glimpse-visa-backend.png")

    add_h(doc, "3.12 Unit testing", 2, "tests")
    add_author(doc, "Written by Krishna Sapkota (S395795). I run this before I call a feature finished. Dharma Bahadur Budhathoki, S395172.")
    add_code(doc, "dotnet test backend\\OraEarn.Api.Tests\\OraEarn.Api.Tests.csproj")
    add_p(
        doc,
        "The white-box tests call the classes directly: shift pay, savings allocation, job checks, the registration review, and the analytics dates. "
        "The black-box tests start the real API on an in-memory database. A bad login is 401. A good login returns a token. Jobs with no token are 401. "
        "The 48-hour block lives in VisaComplianceService and ShiftsController. It runs before the shift is inserted. A bad login is rejected with 401, and a job with no token is rejected the same way.",
    )
    add_figure(doc, "figure-diagram-swagger-tests.png")

    add_h(doc, "3.13 Class diagram and data model", 2, "model")
    add_author(doc, "Written by Suraj Khadka (S394949), with Krishna Sapkota. This is the shape I code against. Dharma Bahadur Budhathoki, S395172.")
    add_p(
        doc,
        "The tables are in third normal form. A shift does not store the worker or the hourly rate. "
        "An expense does not store the worker, and it does not store the words \"Card\" or \"Cash\". It points at a payment method. "
        "TotalHours and TotalEarned stay on the shift, so if I change the job rate later, old payslips stay as they were.",
    )
    add_code(
        doc,
        "ApplicationUser 1 -- * Job 1 -- * Shift\n"
        "        |                 |\n"
        "        |                 +-- * ActiveShiftSession   (one live clock-in per user)\n"
        "        +-- * ExpenseCategory 1 -- * Expense * -- 1 PaymentMethod\n"
        "        +-- * SavingsGoal\n"
        "        +-- * UserNotificationRead",
    )
    add_p(
        doc,
        "The visa file sits on disk. The user row only stores the path and the original file name. "
        "Krishna's write-up of how we got from a spreadsheet to these tables is docs/OraEarn-Normalization-1NF-to-3NF.md.",
    )
    add_figure(doc, "figure-diagram-team-erd.png")
    add_figure(doc, "figure-diagram-uml-classes.png")

    add_h(doc, "4. Adding a new feature", 1, "feature")
    add_author(doc, "Written by Dharma Bahadur Budhathoki (S395172)")
    add_p(
        doc,
        "I'll use a leave request as the example. It has a date, a reason, and a status. Follow the same order I use. Don't start in the controller and hope the rest shows up.",
    )
    add_p(doc, "I make the Jira card first, on the EARN board, something like EARN-### Leave requests. One sentence says what a signed-in user can do. Then the branch:")
    add_code(doc, "git checkout main\ngit pull\ngit checkout -b OE-###-leave-requests")
    add_p(doc, "Backend folder backend/OraEarn.Api/Leave/:")
    add_bullets(
        doc,
        [
            "LeaveRequest.cs holds Id, UserId, Date, Reason, and the link to ApplicationUser.",
            "LeaveDtos.cs holds the request and the response.",
            "LeaveService.cs filters every query by the user id on the token. An empty reason throws ArgumentException.",
            "LeaveController.cs has [Authorize]. It catches ArgumentException and returns BadRequest with { message }.",
        ],
    )
    add_p(doc, "Add the DbSet on ApplicationDbContext, index UserId, and delete the rows with the user, the same way a savings goal works. Then register the service. This is the line I miss when I go too fast:")
    add_code(doc, "services.AddScoped<LeaveService>();")
    add_p(doc, "From backend/OraEarn.Api:")
    add_code(doc, "dotnet ef migrations add AddLeaveRequests --project OraEarn.Api.csproj")
    add_p(doc, "Read the Up method. Restart the API so the migration runs. Only add a NuGet package if you really need one:")
    add_code(doc, "dotnet add package <PackageName>")
    add_p(
        doc,
        "On the website I add frontend/src/features/leave/ with api.ts, LeavePage.tsx, and a dialog. "
        "api.ts uses the shared axios client. I add the route inside ProtectedRoute in App.tsx, and a nav item in Layout.tsx. "
        "I copy the loading text and the empty state from the jobs page.",
    )
    add_p(doc, "If the screen needs a JavaScript package we don't already have, install it from the frontend folder and commit package.json with package-lock.json:")
    add_code(doc, "npm install <package>")
    add_p(
        doc,
        "I don't add another CSS framework. The app already uses Tailwind. I reuse flex, gap-3, rounded-lg, bg-card, and text-muted-foreground. "
        "The colours live in frontend/src/index.css. I only change those variables when the whole app should change, not for one screen.",
    )
    add_p(doc, "The checks in the browser use shared/form-field, the same way the job dialog does. If validate comes back false, I don't post.")
    add_code(
        doc,
        "validate(() => pickErrors([\n"
        "  ['reason', validators.required(form.reason, 'Reason is required.')],\n"
        "  ['date', validators.dateRequired(form.date)],\n"
        "]))",
    )
    add_p(
        doc,
        "The red line under the box is from the browser. The summary at the top is the message from the API, for a rule only the server knows, like the 48-hour limit. "
        "I still keep that check in the service.",
    )
    add_p(doc, "I add one test that an empty reason throws, and one test that GET /api/leave with no token is 401. Then I run dotnet test, and I commit only the leave files:")
    add_code(
        doc,
        'git commit -m "Add leave requests so a user can record time away from shifts."',
    )
    add_p(doc, "I only comment a line when the next person would otherwise guess, like the visa rule or why a delete is blocked. I don't comment \"get the user\". The card goes to Done after someone else has looked at the branch and the tests pass.")
    add_figure(doc, "figure-jobs-desktop.png")
    add_figure(doc, "figure-jobs-add-desktop.png")
    add_figure(doc, "figure-jobs-mobile.png")

    add_h(doc, "5. My contribution", 1, "mine")
    add_author(doc, "Written by Dharma Bahadur Budhathoki (S395172)")
    add_p(
        doc,
        "My name is Dharma Bahadur Budhathoki and my student number is S395172. I am the developer. "
        "The screens a student uses, the sign-in, and the way the API is put together are the work I led. "
        "Krishna led the tables and the test project. Jenish led the requirements. Suraj led the numbers on the charts.",
    )
    add_p(doc, "These are the developer cards I have finished. The key is a link to that card on Jira.")
    # real hyperlinks under the table, because table cells are plain text
    add_table(
        doc,
        ["Key", "Card", "What I delivered"],
        [
            ["EARN-105", "Login", "Email and password sign-in, access and refresh tokens, show-password on the login form."],
            ["EARN-106", "User dashboard", "Personal totals for net, earnings, and expenses, with the period filter."],
            ["EARN-107", "Jobs", "Create, list, update, and delete. The hourly rate on the job is what shift pay uses."],
            ["EARN-70", "Log a shift", "Shift form, calendar, hours and pay calculated on the server."],
        ],
    )
    for key, label in (
        ("EARN-105", "EARN-105 Login"),
        ("EARN-106", "EARN-106 Dashboard"),
        ("EARN-107", "EARN-107 Jobs"),
        ("EARN-70", "EARN-70 Log a shift"),
    ):
        add_mixed(
            doc,
            [
                ("link", label, f"https://students-team-dh8zehye.atlassian.net/browse/{key}"),
            ],
        )
    add_p(
        doc,
        "Register, admin approval, and the Excel and PDF export are in the app. The EARN board screenshot in this handover shows the project cards in Done.",
    )
    add_figure(doc, "figure-shifts-desktop.png")
    add_figure(doc, "figure-dashboard-desktop.png")
    add_figure(doc, "figure-expenses-desktop.png")
    add_figure(doc, "figure-savings-desktop.png")
    add_figure(doc, "figure-timesheet-dharma.png")
    add_p(doc, "Over the weeks I also:")
    add_bullets(
        doc,
        [
            "Set up the React app and the .NET API, including Swagger, PostgreSQL, and Entity Framework.",
            "Built registration, the admin approval screens, expenses, jobs, the shift calendar, the 48-hour bar, and the start of Aura.",
            "Centralised the black-and-white theme, updated the home page, and changed savings so a goal takes a share of leftover pay. That change is the migration AllocateSavingsFromNetMoney.",
            "Wrote this handover, and my timesheet through week 12. The file is under my name, Dharma Bahadur Budhathoki (S395172), weeks 4 to 12. Week 4 is 5.5 hours. Weeks 5 to 12 are 6 hours each. The total is 53.5.",
        ],
    )
    add_mixed(
        doc,
        [
            ("text", "The board I keep current is "),
            ("link", "the EARN Jira board", JIRA),
            ("text", ". The code is in "),
            ("link", "PRT_681_SYDNEY_GROUP-4", GIT),
            ("text", "."),
        ],
    )

    add_page_break(doc)
    add_h(doc, "References", 1, "references")
    add_author(doc, "Written by Dharma Bahadur Budhathoki (S395172). Each title opens the source.")
    add_p(
        doc,
        "I have listed the sources a new developer needs. The web addresses are links. Click the title.",
    )

    refs = [
        ("Group 4 (2026) PRT_681_SYDNEY_GROUP-4. GitHub repository.", GIT),
        (
            "Group 4 (2026) EARN board. Jira, Atlassian.",
            JIRA,
        ),
        (
            "Budhathoki, D.B. (2026) OraEarn handover, Assessment 4. Charles Darwin University.",
            GIT + "/blob/main/oraEarn-WebApp/docs/PRT681_Assessment4_Dharma_Bahadur_Budhathoki_S395172_OraEarn_Handover.docx",
        ),
        (
            "Gurung, J. (2026) OraEarn requirements.",
            GIT + "/blob/main/oraEarn-WebApp/docs/OraEarn-Requirements.md",
        ),
        (
            "Group 4 (2026) OraEarn project proposal.",
            GIT + "/blob/main/oraEarn-WebApp/docs/OraEarn-Project-Proposal.md",
        ),
        (
            "Sapkota, K. (2026) OraEarn database normalisation, UNF to 3NF.",
            GIT + "/blob/main/oraEarn-WebApp/docs/OraEarn-Normalization-1NF-to-3NF.md",
        ),
        (
            "Budhathoki, D.B. (2026) Individual timesheet, weeks 4 to 12.",
            GIT + "/blob/main/(s395172)%20Dharma_Bahadur_Budhathoki",
        ),
        (
            "Microsoft (2026) ASP.NET Core documentation.",
            "https://learn.microsoft.com/en-us/aspnet/core/",
        ),
        (
            "Microsoft (2026) Entity Framework Core.",
            "https://learn.microsoft.com/en-us/ef/core/",
        ),
    ]
    for index, (label, url) in enumerate(refs, start=1):
        p = doc.add_paragraph()
        style_body(p)
        p.paragraph_format.left_indent = Cm(0.75)
        p.paragraph_format.first_line_indent = Cm(-0.75)
        num = p.add_run(f"[{index}]  ")
        _set_run_font(num, "Calibri", 11)
        add_hyperlink(p, label, url)

    core = doc.core_properties
    core.author = "Dharma Bahadur Budhathoki"
    core.title = "PRT681 Assessment 4 OraEarn Handover"
    core.subject = "Dharma Bahadur Budhathoki S395172"
    core.category = "Assessment 4"

    enable_update_fields(doc)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    personal = Path(
        r"C:\projects\clz\oraEarn\(s395172) Dharma_Bahadur_Budhathoki"
        r"\PRT681_Assessment4_Dharma_Bahadur_Budhathoki_S395172_OraEarn_Handover.docx"
    )
    try:
        doc.save(personal)
        print(personal)
    except PermissionError:
        print("personal copy locked; docs copy saved")
    print(OUT)


if __name__ == "__main__":
    build()
